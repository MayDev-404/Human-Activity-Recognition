"""Build the Phase 3 final report (Report Template Part C) from files in the repository.

    python report/build_final_report.py            # final build; refuses placeholders and missing inputs
    python report/build_final_report.py --draft    # shows every problem as a visible [MISSING: ...] marker
    python report/build_final_report.py --root DIR --no-pdf   # read inputs from another tree (dummy inputs)

Inputs (PHASE3_SPEC 6): the official template, report/final/meta.yaml, configs/team.yaml,
report/interim/members.local.yaml (gitignored: registration numbers and emails),
report/final/sections/*.md, report/final/references.yaml, report/final/contributions/<member>.md,
results/final_summary.csv, results/tuning/<model>/tuning_summary.csv, results/analysis/ and the
figures. The template's cover is filled in; everything after it is replaced by the paper in the C.4
format, ending with the C.3 contribution page. Every result number comes from a placeholder and every
paper from a citation key (report/final/README.md). Output: report/build/ICT4442_Final_Report_HAR.docx
(+ .pdf via LibreOffice, or Word on Windows); report/build/ is gitignored because the cover carries
registration numbers.
"""

from __future__ import annotations

import argparse
import copy
import csv
import json
import re
import sys
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_interim_report as interim  # noqa: E402  (shared helpers: markdown, Cursor, tables, members, PDF)

from har.eval.plots import display_name  # noqa: E402
from har.utils.paths import REPO_ROOT  # noqa: E402

BuildError = interim.BuildError
Inputs = interim.Inputs

TEMPLATE = interim.TEMPLATE
OUTPUT_STEM = "ICT4442_Final_Report_HAR"
FINAL_DIR = Path("report") / "final"
SECTION_DIR = FINAL_DIR / "sections"
LOCAL_MEMBERS = Path("report") / "interim" / "members.local.yaml"

# C.4 formatting
FONT, EMAIL_FONT = "Times New Roman", "Courier New"
TITLE_PT, AUTHOR_PT, AFFILIATION_PT, EMAIL_PT = 24, 11, 10, 9
BODY_PT, CAPTION_PT, TABLE_PT, REFERENCE_PT = 10, 8, 8, 8
FIGURE_WIDTH_IN = 3.3
A4_TWIPS = {"w": "11906", "h": "16838"}

MIN_REFERENCES = 12          # Part C minimum
MIN_LITERATURE_PAPERS = 10   # papers cited in Section II
CONTRIBUTION_MAX_WORDS = 60
EM_DASH = chr(0x2014)   # U+2014, never allowed in report text (CLAUDE.md section 6)
C3_HEADERS = ["Member Name", "Reg. No.", "Model / Task", "Key Contributions", "Signature"]
META_KEYS = ("team_no", "title", "month_year", "affiliation", "models")

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-z]+)\s*:\s*([^{}]*?)\s*\}\}")
CITATION_RE = re.compile(r"\[\s*(@[A-Za-z0-9_-]+(?:\s*;\s*@[A-Za-z0-9_-]+)*)\s*\]")


# ---------------------------------------------------------------- report outline

@dataclass(frozen=True)
class Part:
    """One section file: its word budget, optional level-2 heading and the tables/figures after its text."""

    file: str
    budget: tuple[int, int]
    heading: str | None = None
    floats: tuple[str, ...] = ()


# Level-1 sections in order (numbered I, II, ... by the builder); level-2 parts are lettered A, B, ...
OUTLINE: list[tuple[str, list[Part]]] = [
    ("Introduction", [Part("1_introduction.md", (300, 400))]),
    ("Literature Review", [Part("2_literature.md", (450, 550))]),
    ("Objectives", [Part("3_objectives.md", (80, 140))]),
    ("Methodology", [
        Part("4a_mlp.md", (200, 280), "Model 1: MLP"),
        Part("4b_cnn1d.md", (200, 280), "Model 2: 1D-CNN"),
        Part("4c_birnn.md", (250, 330), "Model 3: BiLSTM with GRU Ablation"),
        Part("4d_transformer.md", (200, 280), "Model 4: Transformer Encoder"),
    ]),
    ("Experimental Setup", [Part("5_setup.md", (350, 450), floats=("fig:class_distribution",))]),
    ("Results and Discussion", [
        Part("6a_results.md", (400, 500), "Comparison of the Models",
             ("table:final", "table:tuning", "fig:final_val_curves", "fig:cost_tradeoff")),
        Part("6b_error_analysis.md", (400, 500), "Error Analysis",
             ("table:per_class", "fig:per_class_f1", "fig:error_overlap", "fig:confusion_best")),
    ]),
    ("Conclusion", [Part("7_conclusion.md", (150, 250))]),
    ("Future Work", [Part("8_future_work.md", (120, 200))]),
]
ABSTRACT = Part("abstract.md", (150, 250))
KEYWORDS = Part("keywords.md", (5, 8))   # budget in comma-separated terms
ACKNOWLEDGEMENT = Part("acknowledgement.md", (40, 80))
LITERATURE_FILE = "2_literature.md"

TABLE_CAPTIONS = {
    "final": "Final comparison (mean ± sample std over {seeds})",
    "tuning": "Tuning grids (four configurations per model) and selected configurations",
    "per_class": "Test F1 per class (mean ± sample std over {seeds})",
}
FIGURES = {  # name -> (path, caption); confusion_best is chosen from the final table
    "class_distribution": ("docs/figures/class_distribution.png",
                           "Windows per activity class in the train, validation and test splits."),
    "final_val_curves": ("results/figures/final_val_curves.png",
                         "Validation macro-F1 per epoch of each model's seed-42 final run."),
    "cost_tradeoff": ("results/figures/cost_tradeoff.png",
                      "Test macro-F1 (mean ± std over {seeds}) against single-process CPU latency at batch "
                      "size 1 (log scale); marker area is proportional to the parameter count."),
    "per_class_f1": ("results/analysis/per_class_f1.png", "Test F1 per class and model (mean ± std over {seeds})."),
    "error_overlap": ("results/analysis/error_overlap.png",
                      "Pairwise Jaccard overlap of the misclassified test windows between models, averaged "
                      "over {seeds}."),
    "confusion_best": ("results/figures/confusion_test_{model}_final.png",
                       "Test confusion matrix of the {label}, the model with the best mean test macro-F1 "
                       "(summed over {seeds}, row-normalised)."),
}

# {{final:<model>:<field>}}: (format spec, scale, suffix), as in results/final_summary.md
FINAL_FORMATS = {
    "val_acc": (".1f", 100, "%"), "test_acc": (".1f", 100, "%"),
    "val_macro_f1": (".3f", 1, ""), "test_macro_f1": (".3f", 1, ""),
    "train_time_s": (".1f", 1, ""), "cpu_latency_bs1_ms": (".3f", 1, ""),
}
TUNING_FORMATS = {"val_acc": (".1f", 100, "%"), "val_macro_f1": (".3f", 1, "")}


def roman(n: int) -> str:
    """Upper-case Roman numeral for 1 <= n < 40."""
    numerals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for value, letters in numerals:
        while n >= value:
            out, n = out + letters, n - value
    return out


def float_numbers() -> dict[str, str]:
    """``table:final -> Table I``, ``fig:cost_tradeoff -> Fig. 3``: numbered in order of appearance."""
    numbers, tables, figures = {}, 0, 0
    for _, parts in OUTLINE:
        for part in parts:
            for name in part.floats:
                if name.startswith("table:"):
                    tables += 1
                    numbers[name] = f"Table {roman(tables)}"
                else:
                    figures += 1
                    numbers[name] = f"Fig. {figures}"
    return numbers


def missing(inputs: Inputs, message: str, marker: str) -> str:
    """Report a problem (fatal unless --draft) and return the visible draft marker."""
    inputs.problem(message)
    return f"[MISSING: {marker}]"


def word_count(text: str) -> int:
    """Words of section text as read: citations removed, each placeholder counted as one word."""
    text = PLACEHOLDER_RE.sub("N", CITATION_RE.sub("", interim.plain(text)))
    return len(text.split())


def _float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def format_value(mean: float | None, std: float | None, fmt: tuple[str, float, str], part: str | None) -> str | None:
    """``0.941 ± 0.004``; ``part`` ``mean`` or ``std`` picks one value. None when the value is missing."""
    spec, scale, suffix = fmt
    if part == "std":
        return None if std is None else format(scale * std, spec) + suffix
    if mean is None:
        return None
    text = format(scale * mean, spec)
    if part is None and std is not None:
        text += f" ± {format(scale * std, spec)}"
    return text + suffix


# ---------------------------------------------------------------- inputs

def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """``(column names, rows)`` of a CSV file, values as strings."""
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader.fieldnames or []), list(reader)


def load_final_meta(inputs: Inputs) -> dict[str, Any]:
    """report/final/meta.yaml; any placeholder value is fatal for a final build."""
    meta = interim.load_yaml_file(inputs.path(FINAL_DIR, "meta.yaml")) or {}
    for key in META_KEYS:
        meta.setdefault(key, None)
    unfilled = interim.placeholder_fields(meta)
    if unfilled:
        inputs.problem("report/final/meta.yaml still has placeholder values (empty, FILL_ME, TBD or <...>) in: "
                       + ", ".join(unfilled))
    for key in META_KEYS:
        if key in unfilled:
            meta[key] = [] if key == "models" else f"[MISSING: {key}]"
    return meta


def load_emails(inputs: Inputs, members: list[interim.Member]) -> dict[str, str]:
    """Member id -> email from the ``emails:`` mapping in the gitignored members.local.yaml (D8)."""
    path = inputs.path(LOCAL_MEMBERS)
    local = (interim.load_yaml_file(path) or {}) if path.is_file() else {}
    emails = local.get("emails") or {}
    found = {}
    for member in members:
        email = str(emails.get(member.id, "")).strip()
        if interim.is_placeholder(email) or "@" not in email:
            if path.is_file():  # a missing file was already reported by load_members
                email = missing(inputs, f"no email for {member.id} under 'emails:' in {path}, e.g. "
                                        'emails: {M1: "name@example.com", ...}', f"email of {member.id}")
            else:
                email = f"[MISSING: email of {member.id}]"
        found[member.id] = email
    return found


@dataclass
class ReportData:
    """Everything numbers and tables are read from (PHASE3_SPEC 6.4 and 6.6)."""

    final: dict[str, dict[str, str]] = field(default_factory=dict)        # model -> final_summary.csv row
    tuning: dict[str, tuple[list[str], list[dict[str, str]]]] = field(default_factory=dict)
    analysis: dict[str, Any] = field(default_factory=dict)                # error_summary.json
    per_class: tuple[list[str], list[dict[str, str]]] | None = None       # per_class_f1.csv
    split: dict[str, Any] = field(default_factory=dict)
    seeds: list[int] = field(default_factory=list)
    references: dict[str, str] = field(default_factory=dict)

    @property
    def seeds_text(self) -> str:
        """``seeds 42, 43 and 44``."""
        names = [str(s) for s in self.seeds]
        if len(names) < 2:
            return "seed " + "".join(names) if names else "the final seeds"
        return f"seeds {', '.join(names[:-1])} and {names[-1]}"


def load_data(inputs: Inputs, models: list[str]) -> ReportData:
    """Read every number source; each missing one is a problem (fatal unless --draft)."""
    data = ReportData()
    base = inputs.path("configs", "base.yaml")
    data.seeds = list((interim.load_yaml_file(base) or {}).get("final_seeds") or [])
    if not data.seeds:
        inputs.problem(f"{base} has no final_seeds")
    split = inputs.path("configs", "split.json")
    if split.is_file():
        data.split = json.loads(split.read_text(encoding="utf-8"))
    else:
        inputs.problem(f"{split} not found")
    data.references = {k: str(v).strip() for k, v in
                       (interim.load_yaml_file(inputs.path(FINAL_DIR, "references.yaml")) or {}).items()}

    final = inputs.path("results", "final_summary.csv")
    if final.is_file():
        data.final = {row["model"]: row for row in read_csv(final)[1]}
        for model in models:
            if model not in data.final:
                inputs.problem(f"{final} has no row for {model}; M2 reruns python scripts/make_results_table.py")
    else:
        inputs.problem(f"{final} not found; it is written by python scripts/make_results_table.py (F2.3)")
    for model in models:
        path = inputs.path("results", "tuning", model, "tuning_summary.csv")
        if path.is_file():
            data.tuning[model] = read_csv(path)
        else:
            inputs.problem(f"{path} not found; the model owner runs python scripts/tune.py --model {model}")
    summary = inputs.path("results", "analysis", "error_summary.json")
    if summary.is_file():
        data.analysis = json.loads(summary.read_text(encoding="utf-8"))
    else:
        inputs.problem(f"{summary} not found; it is written by python scripts/error_analysis.py (F3.2)")
    per_class = inputs.path("results", "analysis", "per_class_f1.csv")
    if per_class.is_file():
        data.per_class = read_csv(per_class)
    else:
        inputs.problem(f"{per_class} not found; it is written by python scripts/error_analysis.py (F3.2)")
    return data


def selected_tuning_row(rows: list[dict[str, str]]) -> dict[str, str] | None:
    """The row of a tuning_summary.csv marked ``selected`` (the configuration in configs/final/)."""
    return next((r for r in rows if str(r.get("selected", "")).strip().lower() == "true"), None)


def tuning_parameters(columns: list[str]) -> list[str]:
    """The grid parameter columns of a tuning_summary.csv: those between ``config_id`` and ``params``."""
    if "config_id" in columns and "params" in columns:
        return columns[columns.index("config_id") + 1: columns.index("params")]
    return []


def per_class_cells(table: tuple[list[str], list[dict[str, str]]], models: list[str]
                    ) -> tuple[list[str], dict[tuple[str, str], tuple[float | None, float | None]]]:
    """Classes in file order and ``(class, model) -> (mean, std)`` from per_class_f1.csv.

    Long format: columns ``model``, ``class`` and one column ending in ``mean`` and one in ``std``.
    Wide format: a ``class`` column plus ``<model>_mean`` and ``<model>_std`` columns.
    """
    columns, rows = table
    class_col = next((c for c in columns if c.lower() in ("class", "activity", "label")), None)
    if class_col is None:
        raise ValueError("per_class_f1.csv needs a class column")
    classes = list(dict.fromkeys(r[class_col] for r in rows))
    cells: dict[tuple[str, str], tuple[float | None, float | None]] = {}
    if "model" in columns:
        mean_col = next(c for c in columns if c.lower().endswith("mean"))
        std_col = next(c for c in columns if c.lower().endswith("std"))
        for r in rows:
            cells[(r[class_col], r["model"])] = (_float(r[mean_col]), _float(r[std_col]))
    else:
        for r in rows:
            for model in models:
                if f"{model}_mean" in r:
                    cells[(r[class_col], model)] = (_float(r[f"{model}_mean"]), _float(r.get(f"{model}_std")))
    return classes, cells


# ---------------------------------------------------------------- text: placeholders and citations

class Resolver:
    """Fills ``{{kind:...}}`` placeholders and ``[@key]`` citations, numbering references by first citation."""

    def __init__(self, inputs: Inputs, data: ReportData, numbers: dict[str, str]) -> None:
        self.inputs = inputs
        self.data = data
        self.numbers = numbers
        self.cited: list[str] = []                     # keys in order of first citation
        self.cited_in: dict[str, set[str]] = {}         # section file -> keys cited there

    def text(self, text: str, where: str) -> str:
        """``text`` from the file ``where`` with placeholders and citations filled in (in document order)."""
        text = PLACEHOLDER_RE.sub(lambda m: self.placeholder(m.group(1), m.group(2), where), text)
        text = CITATION_RE.sub(lambda m: self.cite(m.group(1), where), text)
        self.check_dashes(text, where)
        return text

    def check_dashes(self, text: str, where: str) -> None:
        """An em dash in inserted text is a problem (fatal unless --draft)."""
        if EM_DASH in text:
            self.inputs.problem(f"{where} contains an em dash (U+2014); use a comma, colon or parentheses")

    def cite(self, group: str, where: str) -> str:
        """``@a; @b`` -> ``[1], [2]``, numbering each key on its first citation."""
        parts = []
        for key in (k.strip().lstrip("@") for k in group.split(";")):
            if key not in self.data.references:
                parts.append(missing(self.inputs, f"{where}: unknown citation key @{key} "
                                                  "(keys are in report/final/references.yaml)", f"@{key}"))
                continue
            if key not in self.cited:
                self.cited.append(key)
            self.cited_in.setdefault(where, set()).add(key)
            parts.append(f"[{self.cited.index(key) + 1}]")
        return ", ".join(parts)

    def placeholder(self, kind: str, arg: str, where: str) -> str:
        """The value of ``{{kind:arg}}``, or a problem and a ``[MISSING: ...]`` marker."""
        resolvers = {"final": self.final, "tuning": self.tuning, "analysis": self.analysis, "split": self.split,
                     "table": self.float_ref("table"), "fig": self.float_ref("fig")}
        value = None
        if kind in resolvers:
            try:
                value = resolvers[kind](arg)
            except (KeyError, IndexError, TypeError, ValueError):
                value = None
        if value is None:
            return missing(self.inputs, f"{where}: cannot resolve {{{{{kind}:{arg}}}}} "
                                        "(see report/final/README.md)", f"{kind}:{arg}")
        return value

    def final(self, arg: str) -> str | None:
        """``<model>:<field>[.mean|.std]`` from results/final_summary.csv."""
        model, name = arg.split(":")
        name, _, part = name.partition(".")
        row = self.data.final[model]
        if part not in ("", "mean", "std"):
            return None
        if name == "params":
            return f"{int(float(row['params'])):,}" if not part else None
        fmt = FINAL_FORMATS[name]
        return format_value(_float(row.get(f"{name}_mean")), _float(row.get(f"{name}_std")), fmt, part or None)

    def tuning(self, arg: str) -> str | None:
        """``<model>:<column>`` from the selected row of results/tuning/<model>/tuning_summary.csv."""
        model, name = arg.split(":")
        row = selected_tuning_row(self.data.tuning[model][1])
        if row is None or name not in row:
            return None
        if name == "params":
            return f"{int(float(row['params'])):,}"
        if name in TUNING_FORMATS:
            return format_value(_float(row[name]), None, TUNING_FORMATS[name], None)
        return row[name] or None

    def analysis(self, arg: str) -> str | None:
        """``<dotted.key>[|<format>]`` from results/analysis/error_summary.json (see ``format_analysis``)."""
        key, _, spec = (s.strip() for s in arg.partition("|"))
        value: Any = self.data.analysis
        for name in key.split("."):
            value = value[int(name)] if isinstance(value, list) else value[name]
        return format_analysis(value, spec)

    def split(self, arg: str) -> str | None:
        """``<field>`` from configs/split.json."""
        return format_scalar(self.data.split[arg])

    def float_ref(self, kind: str) -> Callable[[str], str | None]:
        """Resolver for ``{{table:<name>}}`` or ``{{fig:<name>}}``: ``Table I``, ``Fig. 3``."""
        return lambda name: self.numbers.get(f"{kind}:{name}")


def format_scalar(value: Any) -> str | None:
    """Analysis and split values: integers as they are, floats to 3 decimals, lists comma-separated."""
    if isinstance(value, bool) or value is None or isinstance(value, dict):
        return None
    if isinstance(value, list):
        items = [format_scalar(v) for v in value]
        return None if None in items else ", ".join(items)  # type: ignore[arg-type]
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) >= 1 else f"{value:.3f}"
    return str(value)


def format_analysis(value: Any, spec: str = "") -> str | None:
    """An error_summary.json value, optionally with a format after ``|``.

    ``spec`` is ``.0f`` to ``.3f`` (decimals) or ``%`` (a fraction as a percentage with one decimal):
    ``0.978|%`` gives ``97.8%``. A ``{"mean", "std"}`` pair prints as ``45.7 ± 5.1`` (``.3f`` without a
    ``spec``). Without a ``spec`` other values print as ``format_scalar`` does. None for an unknown
    ``spec`` or a value it cannot format.
    """
    if spec == "%":
        fmt: tuple[str, float, str] | None = (".1f", 100, "%")
    else:
        fmt = (spec or ".3f", 1, "") if re.fullmatch(r"(\.[0-3]f)?", spec) else None
    if fmt is None:
        return None
    if isinstance(value, dict) and {"mean", "std"} <= set(value):
        return format_value(_float(value["mean"]), _float(value["std"]), fmt, None)
    if not spec:
        return format_scalar(value)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return format_value(float(value), None, fmt, None)


# ---------------------------------------------------------------- docx: page setup and formatting

def set_outline_level(paragraph: Paragraph, level: int) -> None:
    """Show the heading in Word's navigation pane without using the template's heading styles."""
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), str(level))
    paragraph._p.get_or_add_pPr().append(lvl)


def set_run_font(run: Any, name: str) -> None:
    """Set the run's font for every script (Latin, complex and East Asian)."""
    run.font.name = name
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def set_document_font(doc: DocxDocument) -> None:
    """Times New Roman throughout: the document defaults and the Normal style, with no theme fonts."""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    targets = [doc.styles["Normal"].element.get_or_add_rPr()]
    if defaults is not None and defaults.find(qn("w:rPrDefault")) is not None:
        rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        if rpr is not None:
            targets.append(rpr)
    for rpr in targets:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in list(rfonts.attrib):
            if attr.endswith("Theme"):
                del rfonts.attrib[attr]
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), FONT)


def set_page_setup(doc: DocxDocument) -> None:
    """A4 on the template's final section, with no headers, footers or page numbers."""
    sect = doc.element.body.find(qn("w:sectPr"))
    if sect is None:
        raise BuildError("template has no final w:sectPr")
    size = sect.find(qn("w:pgSz"))
    for key, value in A4_TWIPS.items():
        size.set(qn(f"w:{key}"), value)
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg"):
        for element in sect.findall(qn(tag)):
            sect.remove(element)


def set_borders(table: Table) -> None:
    """Single-line grid borders (the template has no table styles to borrow)."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    tbl_pr.append(borders)


class PaperCursor(interim.Cursor):
    """The interim builder's Cursor with the C.4 paper formats (headings, captions, tables, figures)."""

    def text(self, text: str, *, size: float = BODY_PT, align: Any = WD_ALIGN_PARAGRAPH.JUSTIFY,
             space_after: float = 3) -> Paragraph:
        """A body paragraph: 10 pt, justified."""
        return self.paragraph(text, size=size, align=align, space_after=space_after)

    def run_in(self, label: str, text: str) -> Paragraph:
        """``Abstract:`` style: a bold italic run-in label, then the text."""
        paragraph = self.paragraph(align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)
        run = paragraph.add_run(f"{label}: ")
        run.bold, run.italic, run.font.size = True, True, Pt(BODY_PT)
        interim.add_inline(paragraph, text, size=BODY_PT)
        return paragraph

    def bullet(self, text: str) -> Paragraph:
        """A body-text bullet with a hanging indent."""
        paragraph = self.text("• " + text, space_after=2)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        return paragraph

    def heading(self, text: str, level: int) -> Paragraph:
        """Level 1: small caps, centred (``I. Introduction``). Level 2: italic, left (``A. Model 1: MLP``)."""
        center = level == 1
        paragraph = self.paragraph(align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT,
                                   space_after=4 if center else 3)
        paragraph.paragraph_format.space_before = Pt(8 if center else 6)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        run.font.size = Pt(BODY_PT)
        if center:
            run.font.small_caps = True
        else:
            run.italic = True
        set_outline_level(paragraph, level - 1)
        return paragraph

    def table_caption(self, number: str, text: str) -> Paragraph:
        """Above the table, 8 pt small caps: ``TABLE I. Final comparison``."""
        paragraph = self.paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(f"{number.upper()}. {text}")
        run.font.size, run.font.small_caps = Pt(CAPTION_PT), True
        return paragraph

    def figure_caption(self, number: str, text: str) -> Paragraph:
        """Below the figure, 8 pt: ``Fig. 1. ...``."""
        return self.paragraph(f"{number}. {text}", size=CAPTION_PT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    def figure(self, path: Path) -> Paragraph:
        """A centred picture, about 3.3 in wide, kept with its caption."""
        paragraph = self.paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Inches(FIGURE_WIDTH_IN))
        return paragraph

    def table(self, header: list[str], rows: list[list[str]]) -> Table:
        """A centred grid table in 8 pt with a bold header row; first column left, others centred."""
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        self._place(table._tbl)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_borders(table)
        for r, values in enumerate([header, *rows]):
            for c, value in enumerate(values):
                paragraph = table.rows[r].cells[c].paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(value)
                run.font.size = Pt(TABLE_PT)
                run.bold = r == 0 or None
        self.after_table()
        return table

    def insert_table(self, tbl: Any) -> Table:
        """Place an existing table element (the template's C.3 table)."""
        self._place(tbl)
        self.after_table()
        return Table(tbl, self.doc._body)

    def page_break(self) -> None:
        """Start a new page."""
        self.paragraph(space_after=0).add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- build steps

class FinalReportBuilder:
    """Fills a copy of the template; ``build()`` returns the document."""

    def __init__(self, inputs: Inputs, template: Path) -> None:
        self.inputs = inputs
        self.doc = Document(str(template))
        self.numbers = float_numbers()
        self.warnings: list[str] = []

    def build(self) -> DocxDocument:
        """Cover, title block, Abstract to Acknowledgement, References, then the C.3 page."""
        meta = load_final_meta(self.inputs)
        members = interim.load_members(self.inputs)
        emails = load_emails(self.inputs, members)
        self.labels = {m["key"]: m.get("label") or m["key"] for m in meta["models"]}
        self.data = load_data(self.inputs, list(self.labels))
        self.resolver = Resolver(self.inputs, self.data, self.numbers)

        c3_table = copy.deepcopy(interim.find_table(self.doc, C3_HEADERS)._tbl)
        interim.fill_cover(self.doc, meta["title"], meta["month_year"], members, "Final report on")
        cursor = PaperCursor(self.doc, self.remove_template_body())
        set_page_setup(self.doc)
        set_document_font(self.doc)

        self.title_block(cursor, meta, members, emails)
        self.abstract_and_keywords(cursor)
        for i, (title, parts) in enumerate(OUTLINE, start=1):
            cursor.heading(f"{roman(i)}. {title}", 1)
            letters = iter("ABCDEFGH")
            for part in parts:
                if part.heading:
                    cursor.heading(f"{next(letters)}. {part.heading}", 2)
                self.section(cursor, part)
                for name in part.floats:
                    self.place_float(cursor, name)
        cursor.heading("Acknowledgement", 1)
        self.section(cursor, ACKNOWLEDGEMENT)
        self.check_citations()
        self.reference_list(cursor)
        self.contribution_page(cursor, c3_table, meta, members)
        self.set_core_properties(meta, members)
        return self.doc

    # ---- template surgery and front matter

    def remove_template_body(self) -> Any:
        """Delete Parts A, B and C (everything after the cover) but keep the final ``w:sectPr``;
        return the last cover element, after which the paper is inserted."""
        body = self.doc.element.body
        children = list(body.iterchildren())
        start = next((i for i, el in enumerate(children) if interim.element_text(el).strip().startswith("PART A")),
                     None)
        if start is None:
            raise BuildError("template paragraph starting with 'PART A' not found")
        for element in children[start:]:
            if element.tag != qn("w:sectPr"):
                interim.remove_element(element)
        return children[start - 1]

    def title_block(self, cursor: PaperCursor, meta: dict[str, Any], members: list[interim.Member],
                    emails: dict[str, str]) -> None:
        """Title 24 pt, authors 11 pt, affiliation 10 pt italic, emails 9 pt Courier New, all centred."""
        center = WD_ALIGN_PARAGRAPH.CENTER
        self.resolver.check_dashes(meta["title"], "report/final/meta.yaml title")
        cursor.paragraph(meta["title"], size=TITLE_PT, align=center, space_after=8)
        cursor.paragraph(", ".join(m.name for m in members), size=AUTHOR_PT, align=center, space_after=2)
        affiliation = cursor.paragraph(align=center, space_after=2)
        run = affiliation.add_run(str(meta["affiliation"]))
        run.italic, run.font.size = True, Pt(AFFILIATION_PT)
        mails = cursor.paragraph(align=center, space_after=10)
        run = mails.add_run(", ".join(emails[m.id] for m in members))
        run.font.size = Pt(EMAIL_PT)
        set_run_font(run, EMAIL_FONT)

    def read_part(self, part: Part) -> list[interim.Block] | None:
        """Blocks of one section file, after the budget check; None (with a problem) when missing."""
        path = self.inputs.path(SECTION_DIR, part.file)
        if not path.is_file():
            self.inputs.problem(f"{path} not found")
            return None
        _, body = interim.read_markdown(path)
        blocks = interim.markdown_blocks(body)
        texts = [b.text for b in blocks if b.kind in ("para", "bullet")]
        if part is KEYWORDS:
            count, unit = len([t for t in ", ".join(texts).split(",") if t.strip()]), "terms"
        else:
            count, unit = sum(word_count(t) for t in texts), "words"
        low, high = part.budget
        if not low <= count <= high:
            self.warnings.append(f"{SECTION_DIR / part.file}: {count} {unit}, budget {low} to {high}")
        if any(b.kind == "table" for b in blocks):
            self.inputs.problem(f"{path} contains a markdown table; section files hold prose only "
                                "(the builder generates the tables)")
        return blocks

    def abstract_and_keywords(self, cursor: PaperCursor) -> None:
        """One run-in paragraph each, after the title block."""
        for label, part in (("Abstract", ABSTRACT), ("Keywords", KEYWORDS)):
            blocks = self.read_part(part)
            if blocks is None:
                text = f"[MISSING: {SECTION_DIR / part.file}]"
            else:
                text = self.resolver.text(" ".join(b.text for b in blocks if b.kind in ("para", "bullet")),
                                          part.file)
            cursor.run_in(label, text)

    def section(self, cursor: PaperCursor, part: Part) -> None:
        """Paragraphs and bullets of one section file (its own headings are dropped)."""
        blocks = self.read_part(part)
        if blocks is None:
            cursor.text(f"[MISSING: {SECTION_DIR / part.file}]")
            return
        for block in blocks:
            if block.kind == "para":
                cursor.text(self.resolver.text(block.text, part.file))
            elif block.kind == "bullet":
                cursor.bullet(self.resolver.text(block.text, part.file))

    # ---- tables and figures

    def place_float(self, cursor: PaperCursor, name: str) -> None:
        """A generated table (caption above) or a figure (caption below); a draft marks missing inputs."""
        kind, key = name.split(":")
        number = self.numbers[name]
        if kind == "table":
            cursor.table_caption(number, TABLE_CAPTIONS[key].format(seeds=self.data.seeds_text))
            header, rows, note = getattr(self, f"{key}_table")()
            if rows:
                cursor.table(header, rows)
            if note:
                cursor.text(note, size=CAPTION_PT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
            return
        rel, caption = FIGURES[key]
        fmt = {"seeds": self.data.seeds_text}
        if key == "confusion_best":
            best = self.best_model()
            fmt.update(model=best or "<model>", label=self.labels.get(best or "", best or "best model"))
        path = self.inputs.path(rel.format(**fmt))
        if path.is_file():
            cursor.figure(path)
        else:
            cursor.text(missing(self.inputs, f"{path} not found", rel.format(**fmt)), align=WD_ALIGN_PARAGRAPH.CENTER)
        cursor.figure_caption(number, caption.format(**fmt))

    def best_model(self) -> str | None:
        """The model with the best mean test macro-F1 (only chooses which confusion matrix to show)."""
        scored = [(f, m) for m, row in self.data.final.items() if (f := _float(row.get("test_macro_f1_mean"))) is not None]
        return max(scored)[1] if scored else None

    def final_table(self) -> tuple[list[str], list[list[str]], str]:
        """Table I from results/final_summary.csv: ``(header, rows, note)``."""
        header = ["Model", "Params", "Val Macro-F1", "Test Acc", "Test Macro-F1", "Train time (s)",
                  "CPU latency (ms)"]
        fields = ["params", "val_macro_f1", "test_acc", "test_macro_f1", "train_time_s", "cpu_latency_bs1_ms"]
        if not self.data.final:
            return header, [], "[MISSING: results/final_summary.csv]"
        rows, sources = [], []
        for model, label in self.labels.items():
            if model not in self.data.final:
                rows.append([label, "[MISSING: final runs]"] + [""] * (len(fields) - 1))
                continue
            rows.append([label] + [self.resolver.final(f"{model}:{f}") or "n/a" for f in fields])
            sources.append(self.data.final[model].get("cpu_latency_source") or "")
        sources = list(dict.fromkeys(s for s in sources if s))
        note = f"CPU latency at batch size 1 on one thread: {'; '.join(sources)}." if sources else ""
        return header, rows, note

    def tuning_table(self) -> tuple[list[str], list[list[str]], str]:
        """Table II: each model's grid and selected configuration from its tuning_summary.csv."""
        header = ["Model", "Grid", "Selected", "Params", "Val Macro-F1"]
        rows = []
        for model, label in self.labels.items():
            if model not in self.data.tuning:
                rows.append([label, f"[MISSING: results/tuning/{model}/tuning_summary.csv]", "", "", ""])
                continue
            columns, runs = self.data.tuning[model]
            parameters = tuning_parameters(columns)
            grid = " × ".join(f"{p} {{{', '.join(dict.fromkeys(r[p] for r in runs))}}}" for p in parameters)
            chosen = selected_tuning_row(runs)
            if chosen is None:
                rows.append([label, grid, missing(self.inputs, f"no selected row in results/tuning/{model}/",
                                                  "selected configuration"), "", ""])
                continue
            selected = f"{chosen['config_id']}: " + ", ".join(f"{p} {chosen[p]}" for p in parameters)
            rows.append([label, grid, selected, self.resolver.tuning(f"{model}:params") or "n/a",
                         self.resolver.tuning(f"{model}:val_macro_f1") or "n/a"])
        note = ("Seed 42, validation split only; selection by validation macro-F1, ties within 0.001 to fewer "
                "parameters, then to the Phase 2 default (c1).")
        return header, rows, note

    def per_class_table(self) -> tuple[list[str], list[list[str]], str]:
        """Table III: test F1 per class (rows, ``Walking upstairs``) and model (columns) from
        results/analysis/per_class_f1.csv."""
        header = ["Class"] + list(self.labels.values())
        if self.data.per_class is None:
            return header, [], "[MISSING: results/analysis/per_class_f1.csv]"
        try:
            classes, cells = per_class_cells(self.data.per_class, list(self.labels))
        except (StopIteration, ValueError):
            return header, [], missing(self.inputs, "results/analysis/per_class_f1.csv has an unknown layout "
                                                    "(see report/final/README.md)", "per_class_f1.csv layout")
        rows = [[display_name(name)] + [format_value(*cells.get((name, model), (None, None)), (".3f", 1, ""), None)
                                        or "n/a" for model in self.labels] for name in classes]
        return header, rows, ""

    # ---- references and C.3

    def check_citations(self) -> None:
        """At least 12 references cited overall and at least 10 papers in Section II."""
        cited = len(self.resolver.cited)
        if cited < MIN_REFERENCES:
            self.inputs.problem(f"only {cited} references are cited; Part C needs at least {MIN_REFERENCES}")
        literature = len(self.resolver.cited_in.get(LITERATURE_FILE, set()))
        if literature < MIN_LITERATURE_PAPERS:
            self.inputs.problem(f"Section II ({LITERATURE_FILE}) cites {literature} papers; it needs at least "
                                f"{MIN_LITERATURE_PAPERS}")

    def reference_list(self, cursor: PaperCursor) -> None:
        """The cited references, numbered in order of first citation, IEEE style."""
        cursor.heading("References", 1)
        for n, key in enumerate(self.resolver.cited, start=1):
            paragraph = cursor.text(f"[{n}] {self.data.references[key]}", size=REFERENCE_PT, space_after=2)
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.3)

    def contribution_page(self, cursor: PaperCursor, c3_table: Any, meta: dict[str, Any],
                          members: list[interim.Member]) -> None:
        """C.3 on its own last page: the template's table, one row per member, Signature left blank."""
        cursor.page_break()
        cursor.heading("Individual Contribution Statement", 1)
        cursor.text(f"Team {meta['team_no']}: {meta['title']}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
        rows = []
        for member in members:
            model_task, text = self.contribution(member)
            rows.append([member.name, member.reg_no, model_task, text, ""])
        interim.fill_template_table(cursor.insert_table(c3_table), rows, None)

    def contribution(self, member: interim.Member) -> tuple[str, str]:
        """``(model_task, key contributions)`` from report/final/contributions/<member>.md."""
        rel = FINAL_DIR / "contributions" / f"{member.key}.md"
        path = self.inputs.path(rel)
        if not path.is_file():
            marker = missing(self.inputs, f"{path} not found (written by {member.name} personally, 6.5)", str(rel))
            return marker, marker
        front, body = interim.read_markdown(path)
        text = " ".join(b.text for b in interim.markdown_blocks(body) if b.kind in ("para", "bullet"))
        text = interim.plain(text)
        model_task = str(front.get("model_task") or "").strip()
        if interim.is_placeholder(model_task):
            model_task = missing(self.inputs, f"{path}: front matter needs model_task", "model_task")
        if interim.is_placeholder(text):
            text = missing(self.inputs, f"{path}: no contribution text", "key contributions")
        if len(text.split()) > CONTRIBUTION_MAX_WORDS:
            self.warnings.append(f"{rel}: {len(text.split())} words, at most {CONTRIBUTION_MAX_WORDS}")
        self.resolver.check_dashes(model_task + " " + text, str(rel))
        return model_task, text

    def set_core_properties(self, meta: dict[str, Any], members: list[interim.Member]) -> None:
        """Title and authors (the three members) instead of the template's metadata."""
        props = self.doc.core_properties
        props.title = str(meta["title"])
        props.author = props.last_modified_by = "; ".join(m.name for m in members)
        props.subject = "ICT 4442 Deep Learning mini project: final report (Part C)"
        props.keywords = props.comments = props.category = props.content_status = ""
        props.created = props.modified = datetime.now().replace(microsecond=0)
        props.revision = 1


# ---------------------------------------------------------------- output

def build(inputs: Inputs, template: Path, out_dir: Path, make_pdf: bool) -> tuple[Path, list[str]]:
    """Build the docx (and PDF if possible); return its path and the budget warnings."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{OUTPUT_STEM}{'_DRAFT' if inputs.draft else ''}.docx"
    builder = FinalReportBuilder(inputs, template)
    builder.build().save(str(docx_path))
    pdf = interim.convert_to_pdf(docx_path) if make_pdf else None
    if pdf is not None:
        pages = interim.count_pages(pdf)
        print(f"PDF: {pdf} ({pages} pages: the cover plus {pages - 1} pages including C.3; aim for about 7 to 9)")
    if make_pdf and pdf is None:
        print("No PDF converter worked (LibreOffice or Word): open the docx in Word, check it and export the PDF.")
    return docx_path, builder.warnings


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    """Command-line options (see the module docstring)."""
    parser = argparse.ArgumentParser(description="Build the ICT 4442 final report (Part C) from repo files.")
    parser.add_argument("--root", type=Path, default=REPO_ROOT, help="tree to read inputs from (default: repo root)")
    parser.add_argument("--template", type=Path, default=TEMPLATE, help="official report template docx")
    parser.add_argument("--out-dir", type=Path, default=None, help="output folder (default: <root>/report/build)")
    parser.add_argument("--draft", action="store_true",
                        help="show problems as [MISSING: ...] markers instead of failing (writes ..._DRAFT.docx)")
    parser.add_argument("--no-pdf", action="store_true", help="skip PDF conversion")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> Path:
    """Build the report and return the docx path."""
    args = parse_args(argv)
    inputs = Inputs(root=args.root.resolve(), draft=args.draft)
    out_dir = args.out_dir or inputs.path("report", "build")
    docx_path, warnings = build(inputs, args.template, out_dir, make_pdf=not args.no_pdf)
    for message in warnings:
        print(f"WARNING: {message}")
    for message in inputs.problems:
        print(f"DRAFT WARNING: {message}")
    print(f"Report written to {docx_path}. Open it in Word and check every number against results/final_summary.csv "
          "and results/analysis/.")
    return docx_path


if __name__ == "__main__":
    try:
        main()
    except BuildError as exc:
        sys.exit(f"build_final_report: {exc}")
