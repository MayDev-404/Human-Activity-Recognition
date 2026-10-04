"""Build the Phase 2 interim report (Report Template Part B) from files in the repository.

    python report/build_interim_report.py            # final build; refuses placeholders and missing inputs
    python report/build_interim_report.py --draft    # allow FILL_ME, missing reg. nos. and missing inputs
    python report/build_interim_report.py --root DIR --no-pdf   # read inputs from another tree (dummy runs)

Inputs (BUILD_SPEC 10.2 and 10.3): the official template, report/interim/meta.yaml, configs/team.yaml,
report/interim/members.local.yaml (gitignored), docs/lit/<member>.md, report/interim/sections/*.md,
results/summary.csv, docs/tasks/<member>.md, docs/contribution_log.md and two figures. Every number in
the report comes from results/summary.csv. Output: report/build/ICT4442_Interim_Report_HAR.docx
(+ .pdf via LibreOffice, or Word on Windows); report/build/ is gitignored because the cover carries
registration numbers.
"""

from __future__ import annotations

import argparse
import copy
import csv
import json
import math
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from har.utils.paths import REPO_ROOT

TEMPLATE = REPO_ROOT / "report" / "template" / "ICT_4442_Mini_Project_Report_Template.docx"
OUTPUT_STEM = "ICT4442_Interim_Report_HAR"
PLACEHOLDER = "FILL_ME"
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

BODY_PT = 10
TABLE_PT = 9
LIT_TABLE_PT = 8
CAPTION_PT = 8
REFERENCE_PT = 8
FIGURE_WIDTH_IN = 3.2

# Template tables, identified by the start of each header cell (never by index).
LIT_HEADERS = ["Paper", "Method", "Dataset", "Key Result", "Relevance"]
MODEL_HEADERS = ["Model", "Owner", "Status", "Preliminary Metric", "Notes"]
CONTRIB_HEADERS = ["Member Name", "Reg. No.", "Task(s) Completed", "Signature"]

SECTION_DIR = Path("report") / "interim" / "sections"
PREPROCESSING_PARTS = [
    ("preprocessing_raw.md", "Raw-signal pipeline"),
    ("preprocessing_features.md", "Engineered-feature pipeline"),
    ("evaluation_protocol.md", "Shared training and evaluation protocol"),
]


class BuildError(Exception):
    """A required input is missing or invalid; the message says how to fix it."""


# ---------------------------------------------------------------- markdown helpers

COMMENT_RE = re.compile(r"<!--.*?-->", re.S)
FRONT_MATTER_RE = re.compile(r"\A---\s*\n(.*?)\n---\s*(\n|\Z)", re.S)
SEPARATOR_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def read_markdown(path: Path) -> tuple[dict[str, Any], str]:
    """Return ``(front matter, body)`` of a markdown file, with HTML comments removed."""
    text = path.read_text(encoding="utf-8")
    meta: dict[str, Any] = {}
    match = FRONT_MATTER_RE.match(text)
    if match:
        meta = yaml.safe_load(match.group(1)) or {}
        text = text[match.end():]
    return meta, COMMENT_RE.sub("", text).strip()


def split_cells(line: str) -> list[str]:
    """Cells of one markdown table row."""
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_start(lines: list[str], i: int) -> bool:
    return lines[i].strip().startswith("|") and i + 1 < len(lines) and bool(SEPARATOR_RE.match(lines[i + 1].strip()))


def parse_tables(text: str) -> list[tuple[list[str], list[list[str]]]]:
    """Every markdown table in ``text`` as ``(header, rows)``."""
    return [block.table for block in markdown_blocks(text) if block.kind == "table" and block.table]


@dataclass
class Block:
    """One piece of section text: a paragraph, bullet, heading (with level) or table."""

    kind: str
    text: str = ""
    level: int = 0
    table: tuple[list[str], list[list[str]]] | None = None


def markdown_blocks(text: str) -> list[Block]:
    """Split simple markdown into blocks: paragraphs, ``-`` bullets, ``#`` headings and tables."""
    blocks: list[Block] = []
    paragraph: list[str] = []
    lines = text.splitlines()

    def flush() -> None:
        if paragraph:
            blocks.append(Block("para", " ".join(paragraph)))
            paragraph.clear()

    i = 0
    while i < len(lines):
        line, stripped = lines[i], lines[i].strip()
        if not stripped:
            flush()
        elif _is_table_start(lines, i):
            flush()
            header, rows = split_cells(stripped), []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_cells(lines[i]))
                i += 1
            blocks.append(Block("table", table=(header, rows)))
            continue
        elif heading := re.match(r"^(#{1,6})\s+(.*)", stripped):
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), level=len(heading.group(1))))
        elif bullet := re.match(r"^[-*+]\s+(.*)", stripped):
            flush()
            blocks.append(Block("bullet", bullet.group(1).strip()))
        elif line[:1] in (" ", "\t") and blocks and blocks[-1].kind == "bullet" and not paragraph:
            blocks[-1].text += " " + stripped  # wrapped bullet line
        else:
            paragraph.append(stripped)
        i += 1
    flush()
    return blocks


def plain(text: str) -> str:
    """Strip inline markdown (links, bold, italics, code) to plain text."""
    text = LINK_RE.sub(r"\1", text)
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*\s][^*]*)\*|`([^`]+)`", lambda m: next(g for g in m.groups() if g), text)


def add_inline(paragraph: Paragraph, text: str, size: float | None = None, bold: bool | None = None) -> None:
    """Append ``text`` to ``paragraph`` as runs, rendering **bold**, *italic* and `code` spans."""
    for part in INLINE_RE.split(LINK_RE.sub(r"\1", text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if bold is not None and run.bold is None:
            run.bold = bold


# ---------------------------------------------------------------- docx helpers

def element_text(element: Any) -> str:
    """All text inside a body element (paragraph or table)."""
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def remove_element(element: Any) -> None:
    element.getparent().remove(element)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace the paragraph's text in place, keeping run formatting, page breaks and drawings."""
    placed = False
    for run in paragraph.runs:
        for t in run._r.findall(qn("w:t")):
            if not placed:
                t.text = text
                t.set(qn("xml:space"), "preserve")
                placed = True
            else:
                run._r.remove(t)
    if not placed:
        paragraph.add_run(text)


def set_cell_text(cell: _Cell, text: str, size: float | None) -> None:
    """Write ``text`` into the first run of the cell's first paragraph, keeping its formatting."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        remove_element(extra._p)
    runs = paragraph.runs
    if runs:
        set_paragraph_text(paragraph, text)
        target_runs = paragraph.runs[:1]
    else:
        target_runs = [paragraph.add_run(text)]
    if size:
        for run in target_runs:
            run.font.size = Pt(size)


def find_table(doc: DocxDocument, headers: list[str]) -> Table:
    """The template table whose header cells start with ``headers``."""
    for table in doc.tables:
        cells = [c.text.strip().lower() for c in table.rows[0].cells]
        if len(cells) == len(headers) and all(c.startswith(h.lower()) for c, h in zip(cells, headers)):
            return table
    raise BuildError(f"template table with header {headers} not found")


def fill_template_table(table: Table, rows: list[list[str]], size: float | None) -> None:
    """Replace the template's empty rows with ``rows``, deep-copying the first empty row (borders, fonts)."""
    if len(table.rows) < 2:
        raise BuildError("template table has no empty row to copy")
    row_template = copy.deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for values in rows:
        tr = copy.deepcopy(row_template)
        table._tbl.append(tr)
        for tc, value in zip(tr.findall(qn("w:tc")), values):
            set_cell_text(_Cell(tc, table), value, size)


def find_paragraph(doc: DocxDocument, prefix: str) -> Paragraph:
    """The first body paragraph whose stripped text starts with ``prefix``."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise BuildError(f"template paragraph starting with {prefix!r} not found")


class Cursor:
    """Inserts new body content one element after another, starting after ``anchor``."""

    def __init__(self, doc: DocxDocument, anchor: Any, table_style_source: Table) -> None:
        self.doc = doc
        self.element = anchor
        self._tblPr = copy.deepcopy(table_style_source._tbl.tblPr)
        self._header_cell = copy.deepcopy(table_style_source.rows[0].cells[0]._tc)

    def _place(self, element: Any) -> None:
        self.element.addnext(element)
        self.element = element

    def paragraph(self, text: str = "", *, style: str | None = None, size: float | None = BODY_PT,
                  align: Any = None, bold: bool | None = None, space_after: float = 4) -> Paragraph:
        paragraph = Paragraph(OxmlElement("w:p"), self.doc._body)
        self._place(paragraph._p)
        if style:
            paragraph.style = self.doc.styles[style]
        if align is not None:
            paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(space_after)
        if text:
            add_inline(paragraph, text, size=size, bold=bold)
        return paragraph

    def caption(self, text: str) -> None:
        self.paragraph(text, size=CAPTION_PT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)

    def picture(self, path: Path, caption: str) -> None:
        paragraph = self.paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        paragraph.add_run().add_picture(str(path), width=Inches(FIGURE_WIDTH_IN))
        self.caption(caption)

    def table(self, header: list[str], rows: list[list[str]], size: float = TABLE_PT) -> Table:
        """A new table styled like the template tables (borders, dark header row)."""
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        self._place(table._tbl)
        table._tbl.remove(table._tbl.tblPr)
        table._tbl.insert(0, copy.deepcopy(self._tblPr))
        width = int(self._tblPr.find(qn("w:tblW")).get(qn("w:w")))
        for grid_col in table._tbl.tblGrid.findall(qn("w:gridCol")):
            grid_col.set(qn("w:w"), str(width // len(header)))
        for i, value in enumerate(header):
            tc = copy.deepcopy(self._header_cell)
            old = table.rows[0].cells[i]._tc
            old.addnext(tc)
            remove_element(old)
            tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(width // len(header)))
            set_cell_text(_Cell(tc, table), plain(value), size)
        for r, values in enumerate(rows, start=1):
            for c, value in enumerate(values[: len(header)]):
                cell = table.rows[r].cells[c]
                cell._tc.get_or_add_tcPr().get_or_add_tcW().set(qn("w:w"), str(width // len(header)))
                cell.paragraphs[0].text = ""
                add_inline(cell.paragraphs[0], value, size=size)
        return table


# ---------------------------------------------------------------- inputs

@dataclass
class Member:
    id: str
    name: str
    github: str
    reg_no: str

    @property
    def key(self) -> str:
        """File key for docs/lit/<key>.md and docs/tasks/<key>.md (lower-case first name)."""
        return self.name.split()[0].lower()


@dataclass
class LitRow:
    cells: list[str]
    reference: str


@dataclass
class Inputs:
    root: Path
    draft: bool
    problems: list[str] = field(default_factory=list)

    def path(self, *parts: str | Path) -> Path:
        return self.root.joinpath(*parts)

    def problem(self, message: str) -> None:
        """A missing or incomplete input: fatal for a final build, a warning for a draft."""
        if not self.draft:
            raise BuildError(message + " (pass --draft to build anyway)")
        self.problems.append(message)


def load_yaml_file(path: Path) -> Any:
    if not path.is_file():
        raise BuildError(f"{path} not found")
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_members(inputs: Inputs) -> list[Member]:
    team = load_yaml_file(inputs.path("configs", "team.yaml"))["members"]
    reg_path = inputs.path("report", "interim", "members.local.yaml")
    reg_nos: dict[str, Any] = {}
    if reg_path.is_file():
        reg_nos = load_yaml_file(reg_path) or {}
    else:
        inputs.problem(f"{reg_path} not found; create it (gitignored) with the registration numbers from "
                       'the synopsis, e.g. {M1: "<reg no>", M2: "<reg no>", M3: "<reg no>"}')
    members = []
    for entry in team:
        reg = str(reg_nos.get(entry["id"], "")).strip()
        if not reg:
            if reg_path.is_file():
                inputs.problem(f"no registration number for {entry['id']} in {reg_path}")
            reg = "Reg. No. TBD"
        members.append(Member(entry["id"], entry["name"], entry.get("github") or "", reg))
    return members


def load_meta(inputs: Inputs) -> dict[str, Any]:
    meta = load_yaml_file(inputs.path("report", "interim", "meta.yaml"))
    if str(meta.get("team_no", PLACEHOLDER)).strip() in ("", PLACEHOLDER):
        if not inputs.draft:
            raise BuildError("team_no is still FILL_ME in report/interim/meta.yaml; set it (or pass --draft)")
        meta["team_no"] = "TBD"
    return meta


def https_repo_url(url: str) -> str:
    """Normalise a git remote URL to ``https://host/owner/repo`` (no credentials, no .git)."""
    url = url.strip()
    if m := re.match(r"^git@([^:]+):(.+?)(\.git)?$", url):
        return f"https://{m.group(1)}/{m.group(2)}"
    url = re.sub(r"^(https?://)[^@/]+@", r"\1", url)
    return re.sub(r"\.git$", "", url)


def resolve_repo_url(inputs: Inputs, meta: dict[str, Any]) -> str:
    configured = str(meta.get("repo_url", "auto")).strip()
    if configured and configured != "auto":
        return https_repo_url(configured)
    for directory in (inputs.root, REPO_ROOT):
        try:
            result = subprocess.run(["git", "-C", str(directory), "remote", "get-url", "origin"],
                                    capture_output=True, text=True, check=True)
            return https_repo_url(result.stdout)
        except (OSError, subprocess.CalledProcessError):
            continue
    inputs.problem("repo_url is 'auto' but `git remote get-url origin` failed")
    return "GitHub link TBD"


def load_literature(inputs: Inputs, members: list[Member]) -> list[LitRow]:
    """Literature rows in report order (M1, M2, M3), each paired with its reference entry."""
    rows: list[LitRow] = []
    for member in members:
        path = inputs.path("docs", "lit", f"{member.key}.md")
        if not path.is_file():
            inputs.problem(f"{path} not found (literature rows of {member.name})")
            continue
        _, body = read_markdown(path)
        tables = parse_tables(body)
        table_rows = [r for header, rs in tables[:1] for r in rs if any(cell for cell in r)]
        references_section = re.split(r"^#+\s*References\s*$", body, flags=re.M | re.I)
        references = []
        if len(references_section) > 1:
            references = [b.text for b in markdown_blocks(references_section[1]) if b.kind == "bullet"]
        if len(table_rows) != len(references):
            raise BuildError(f"{path}: {len(table_rows)} table rows but {len(references)} references; "
                             "list one reference per row, in the same order")
        rows += [LitRow([plain(c) for c in r[:5]], ref) for r, ref in zip(table_rows, references)]
    return rows


def load_summary(inputs: Inputs) -> dict[str, dict[str, str]]:
    """results/summary.csv as ``{model_key: {normalised_column: value}}``."""
    path = inputs.path("results", "summary.csv")
    if not path.is_file():
        inputs.problem(f"{path} not found; run python scripts/make_results_table.py")
        return {}
    with path.open(newline="", encoding="utf-8") as f:
        raw_rows = list(csv.DictReader(f))
    summary = {}
    for raw in raw_rows:
        row = {re.sub(r"[^a-z0-9]+", "_", k.strip().lower()).strip("_"): (v or "").strip()
               for k, v in raw.items() if k}
        key = row.get("model_key") or row.get("key") or row.get("model")
        if key:
            summary[key] = row
    return summary


RUN_ID_TIME_RE = re.compile(r"^(\d{8}-\d{6})")
REGENERATE = "M2 must rerun python scripts/make_results_table.py and commit the result"


def full_runs(inputs: Inputs) -> dict[str, list[Path]]:
    """Full run folders per model (``results/<model>/<run_id>/metrics.json``).

    Smoke runs (folders starting with ``_``) and protocol-override runs are skipped: they never
    get a row in results/summary.csv.
    """
    runs: dict[str, list[Path]] = {}
    for metrics in sorted(inputs.path("results").glob("*/*/metrics.json")):
        model_dir = metrics.parent.parent
        if model_dir.name.startswith("_"):
            continue
        try:
            if json.loads(metrics.read_text(encoding="utf-8")).get("protocol_override"):
                continue
        except (OSError, ValueError):
            pass  # an unreadable metrics.json still counts as a run that needs a summary row
        runs.setdefault(model_dir.name, []).append(metrics.parent)
    return runs


def run_started(run_dir: Path) -> datetime | None:
    """Start time encoded in a run id such as ``20261004-194905_134ae8c`` (local time)."""
    match = RUN_ID_TIME_RE.match(run_dir.name)
    return datetime.strptime(match.group(1), "%Y%m%d-%H%M%S") if match else None


def summary_written_at(path: Path) -> datetime:
    """When results/summary.csv was last regenerated: the time of its last git commit, or its
    file time if it has uncommitted changes or git cannot tell (file times change on checkout)."""
    def git(*args: str) -> str:
        return subprocess.run(["git", "-C", str(path.parent), *args, "--", path.name],
                              capture_output=True, text=True, check=True).stdout.strip()

    try:
        committed, dirty = git("log", "-1", "--format=%ct"), git("status", "--porcelain")
        if committed and not dirty:
            return datetime.fromtimestamp(int(committed))
    except (OSError, subprocess.CalledProcessError, ValueError):
        pass
    return datetime.fromtimestamp(path.stat().st_mtime)


def check_summary_is_current(inputs: Inputs, summary: dict[str, dict[str, str]]) -> None:
    """Refuse a stale results table: every model with a full run folder needs a summary row, and
    summary.csv must have been regenerated after the newest run started."""
    runs = full_runs(inputs)
    for model, run_dirs in sorted(runs.items()):
        if model not in summary:
            inputs.problem(f"results/{model}/ has {len(run_dirs)} full run folder(s) but results/summary.csv "
                           f"has no '{model}' row; {REGENERATE}")
    path = inputs.path("results", "summary.csv")
    started = [(t, d) for dirs in runs.values() for d in dirs if (t := run_started(d)) is not None]
    if not path.is_file() or not started:
        return
    written = summary_written_at(path)
    newest_time, newest = max(started)
    if newest_time > written:
        inputs.problem(f"results/summary.csv (regenerated {written:%Y-%m-%d %H:%M:%S}) is older than the newest "
                       f"run folder results/{newest.parent.name}/{newest.name}; {REGENERATE}")


def _number(row: dict[str, str], *names: str) -> float | None:
    for name in names:
        value = row.get(name, "")
        try:
            number = float(value)
        except ValueError:
            continue
        if not math.isnan(number):
            return number
    return None


def preliminary_metric(row: dict[str, str] | None) -> str:
    """``Test acc XX.X% / macro-F1 0.XXX (val macro-F1 0.XXX)`` from one summary row."""
    if row is None:
        return "Not trained yet"
    test_acc = _number(row, "test_acc", "test_accuracy")
    test_f1 = _number(row, "test_macro_f1")
    val_f1 = _number(row, "val_macro_f1")
    val_text = f"val macro-F1 {val_f1:.3f}" if val_f1 is not None else "val macro-F1 n/a"
    if test_acc is None or test_f1 is None:
        return f"Test not evaluated ({val_text})"
    if test_acc <= 1.0:
        test_acc *= 100
    return f"Test acc {test_acc:.1f}% / macro-F1 {test_f1:.3f} ({val_text})"


def load_model_notes(inputs: Inputs, key: str) -> tuple[dict[str, Any], str]:
    """Front matter and summary text for a model; ablations (e.g. gru) may live in another file."""
    path = inputs.path(SECTION_DIR, f"model_{key}.md")
    if path.is_file():
        return read_markdown(path)
    for other in sorted(inputs.path(SECTION_DIR).glob("model_*.md")):
        front, _ = read_markdown(other)
        ablation = (front.get("ablations") or {}).get(key)
        if ablation:
            return ablation, ""
    return {}, ""


def done_tasks(inputs: Inputs, member: Member) -> str:
    path = inputs.path("docs", "tasks", f"{member.key}.md")
    if not path.is_file():
        inputs.problem(f"{path} not found")
        return ""
    _, body = read_markdown(path)
    tables = parse_tables(body)
    if not tables:
        return ""
    header, rows = tables[0]
    columns = [h.lower() for h in header]
    status_i, task_i, id_i = columns.index("status"), columns.index("task"), columns.index("id")
    return "; ".join(f"{r[id_i]} {plain(r[task_i])}" for r in rows if r[status_i].strip().lower() == "done")


# ---------------------------------------------------------------- build steps

class ReportBuilder:
    """Fills a copy of the template; ``build()`` returns the document."""

    def __init__(self, inputs: Inputs, template: Path, include_references: bool = True) -> None:
        self.inputs = inputs
        self.include_references = include_references
        self.doc = Document(str(template))
        self.table_no = 0
        self.figure_no = 0

    def next_table_caption(self, title: str) -> str:
        self.table_no += 1
        return f"TABLE {ROMAN[self.table_no - 1]}. {title.upper()}"

    def next_figure_caption(self, text: str) -> str:
        self.figure_no += 1
        return f"Fig. {self.figure_no}. {text}"

    def build(self) -> DocxDocument:
        meta = load_meta(self.inputs)
        members = load_members(self.inputs)
        repo_url = resolve_repo_url(self.inputs, meta)
        summary = load_summary(self.inputs)
        check_summary_is_current(self.inputs, summary)
        literature = load_literature(self.inputs, members)

        self.fill_cover(meta, members)
        self.remove_part_a_and_c()
        self.fill_part_b_header(meta, members, repo_url)
        self.literature_section(literature)
        self.preprocessing_section()
        self.models_section(meta, members, summary)
        self.contribution_section(members, repo_url)
        self.risk_section()
        if self.include_references:
            self.references(literature, meta)
        return self.doc

    # ---- cover and template surgery

    def fill_cover(self, meta: dict[str, Any], members: list[Member]) -> None:
        set_paragraph_text(find_paragraph(self.doc, "<TITLE OF THE PROJECT"), meta["title"])
        set_paragraph_text(find_paragraph(self.doc, "Synopsis/Interim/Final report on"), "Interim report on")
        name_paragraphs = [p for p in self.doc.paragraphs if p.text.strip().startswith("<NAME OF THE STUDENT-")]
        if len(name_paragraphs) != len(members):
            raise BuildError(f"template has {len(name_paragraphs)} student lines but team.yaml lists {len(members)}")
        for paragraph, member in zip(name_paragraphs, members):
            set_paragraph_text(paragraph, f"{member.name} - {member.reg_no}")
        month = next((p for p in self.doc.paragraphs
                      if p.text.strip().startswith("<") and "MONTH AND YEAR" in p.text), None)
        if month is None:
            raise BuildError("template paragraph for the month and year not found")
        set_paragraph_text(month, meta["month_year"])

    def remove_part_a_and_c(self) -> None:
        body = self.doc.element.body
        children = list(body.iterchildren())
        texts = [element_text(el).strip() for el in children]
        start_a = next(i for i, t in enumerate(texts) if t.startswith("PART A"))
        start_b = next(i for i, t in enumerate(texts) if t.startswith("PART B"))
        start_c = next(i for i, t in enumerate(texts) if t.startswith("PART C"))
        for element in children[start_a:start_b] + children[start_c:]:
            if element.tag != qn("w:sectPr"):  # keep page setup
                remove_element(element)
        # Part B ended with an empty paragraph holding the page break before Part C; drop it so
        # the report does not end on a blank page.
        while True:
            last = [el for el in body.iterchildren() if el.tag != qn("w:sectPr")][-1]
            if last.tag == qn("w:p") and not element_text(last).strip() and not last.findall(".//" + qn("w:drawing")):
                remove_element(last)
            else:
                break

    def fill_part_b_header(self, meta: dict[str, Any], members: list[Member], repo_url: str) -> None:
        set_paragraph_text(find_paragraph(self.doc, "PART B"), "INTERIM REPORT")
        for prefix in ("Interim report should be", "Minimum 8"):
            remove_element(find_paragraph(self.doc, prefix)._p)
        names = "; ".join(f"{m.name} ({m.reg_no})" for m in members)
        values = {
            "Team No. and Names": f"Team {meta['team_no']}: {names}",
            "Title of the Project": f"{meta['title']} ({meta.get('title_status', 'Confirmed from Synopsis')})",
            "GitHub Repository Link": repo_url,
        }
        for prefix, value in values.items():
            paragraph = find_paragraph(self.doc, prefix)
            label_run = paragraph.runs[-1]
            run = paragraph.add_run(" " + value)
            if label_run._r.rPr is not None:
                run._r.insert(0, copy.deepcopy(label_run._r.rPr))
            run.bold = False

    def cursor_after(self, element: Any) -> Cursor:
        return Cursor(self.doc, element, find_table(self.doc, LIT_HEADERS))

    def section_blocks(self, cursor: Cursor, filename: str) -> None:
        """Insert one section file's text at the cursor (missing file: draft warning)."""
        path = self.inputs.path(SECTION_DIR, filename)
        if not path.is_file():
            self.inputs.problem(f"{path} not found")
            return
        _, body = read_markdown(path)
        self.insert_blocks(cursor, markdown_blocks(body))

    def insert_blocks(self, cursor: Cursor, blocks: list[Block]) -> None:
        previous_heading = ""
        for block in blocks:
            if block.kind == "heading":
                if block.level > 1:  # level-1 headings repeat the template's section heading
                    cursor.paragraph(block.text, bold=True, space_after=2)
                    previous_heading = plain(block.text)
                continue
            if block.kind == "para":
                cursor.paragraph(block.text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif block.kind == "bullet":
                cursor.paragraph("\u2022 " + block.text, style="List Paragraph", space_after=2)
            elif block.kind == "table" and block.table:
                cursor.caption(self.next_table_caption(previous_heading or "Table"))
                cursor.table(*block.table)
                cursor.paragraph(space_after=2)
            previous_heading = ""

    # ---- sections

    def literature_section(self, literature: list[LitRow]) -> None:
        table = find_table(self.doc, LIT_HEADERS)
        cursor = self.cursor_after(table._tbl.getprevious())
        cursor.caption(self.next_table_caption("Literature review summary"))
        rows = []
        for number, row in enumerate(literature, start=1):
            cells = (row.cells + [""] * 5)[:5]
            cells[0] = f"{cells[0]} [{number}]"
            rows.append(cells)
        fill_template_table(table, rows, LIT_TABLE_PT)
        cursor = self.cursor_after(table._tbl)
        cursor.paragraph(space_after=2)
        self.section_blocks(cursor, "lit_summary.md")

    def preprocessing_section(self) -> None:
        label = find_paragraph(self.doc, "Description of preprocessing pipeline")
        cursor = self.cursor_after(label._p)
        for filename, title in PREPROCESSING_PARTS:
            cursor.paragraph(title, bold=True, space_after=2)
            self.section_blocks(cursor, filename)
        figure = self.inputs.path("docs", "figures", "class_distribution.png")
        if figure.is_file():
            cursor.picture(figure, self.next_figure_caption(
                "Windows per activity class in the train, validation and test splits."))
        else:
            self.inputs.problem(f"{figure} not found; run python scripts/data_report.py")

    def models_section(self, meta: dict[str, Any], members: list[Member], summary: dict[str, dict[str, str]]) -> None:
        owners = {m.id: m.name for m in members}
        table = find_table(self.doc, MODEL_HEADERS)
        cursor = self.cursor_after(table._tbl.getprevious())
        cursor.caption(self.next_table_caption("Models implemented so far (preliminary, untuned, single seed)"))
        rows, summaries = [], []
        for model in meta["models"]:
            front, body = load_model_notes(self.inputs, model["key"])
            in_summary = summary.get(model["key"])
            status = front.get("status") or ("Implemented, preliminary run" if in_summary else "Pending")
            if not front:
                self.inputs.problem(f"no model notes for {model['key']} "
                                    f"(report/interim/sections/model_{model['key']}.md)")
            rows.append([model["label"], owners.get(model["owner"], model["owner"]), str(status),
                         preliminary_metric(in_summary), plain(str(front.get("notes", "")))])
            if body:
                summaries.append((model["label"], body))
        fill_template_table(table, rows, TABLE_PT)
        cursor = self.cursor_after(table._tbl)
        cursor.paragraph(space_after=2)
        for label, body in summaries:
            text = " ".join(b.text for b in markdown_blocks(body) if b.kind in ("para", "bullet"))
            cursor.paragraph(f"**{label}.** {text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.section_blocks(cursor, "observations.md")
        best = max(((k, _number(r, "val_macro_f1")) for k, r in summary.items()),
                   key=lambda kv: -1 if kv[1] is None else kv[1], default=None)
        if best and best[1] is not None:
            figure = self.inputs.path("results", "figures", f"confusion_test_{best[0]}.png")
            label = next((m["label"] for m in meta["models"] if m["key"] == best[0]), best[0])
            if figure.is_file():
                cursor.picture(figure, self.next_figure_caption(
                    f"Test-set confusion matrix of the {label}, the model with the best validation "
                    "macro-F1 (preliminary, untuned, single seed)."))
            else:
                self.inputs.problem(f"{figure} not found; run python scripts/plot_curves.py")

    def contribution_section(self, members: list[Member], repo_url: str) -> None:
        table = find_table(self.doc, CONTRIB_HEADERS)
        cursor = self.cursor_after(table._tbl.getprevious())
        cursor.caption(self.next_table_caption("Individual contribution log (signed by all members)"))
        fill_template_table(table, [[m.name, m.reg_no, done_tasks(self.inputs, m), ""] for m in members], TABLE_PT)
        cursor = self.cursor_after(table._tbl)
        cursor.paragraph(space_after=2)
        path = self.inputs.path("docs", "contribution_log.md")
        if not path.is_file():
            self.inputs.problem(f"{path} not found; run python scripts/contribution_log.py")
            return
        _, body = read_markdown(path)
        tables = parse_tables(body)
        sha = re.search(r"HEAD\W+([0-9a-f]{7,40})", body)
        at = f" at commit {sha.group(1)[:7]}" if sha else ""
        cursor.paragraph(
            f"Commit evidence: the table below is generated from the git history of the main branch{at} "
            f"by scripts/contribution_log.py (full log in docs/contribution_log.md). Every commit is "
            f"visible at {repo_url}/commits/main.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        if tables:
            cursor.caption(self.next_table_caption("Commit summary per member"))
            cursor.table(*tables[0])

    def risk_section(self) -> None:
        label = find_paragraph(self.doc, "Remaining models to implement")
        self.section_blocks(self.cursor_after(label._p), "risk_plan.md")

    def references(self, literature: list[LitRow], meta: dict[str, Any]) -> None:
        body = self.doc.element.body
        last = [el for el in body.iterchildren() if el.tag != qn("w:sectPr")][-1]
        cursor = self.cursor_after(last)
        cursor.paragraph("References", style="Heading 2", size=None, space_after=4)
        entries = [row.reference for row in literature]
        if meta.get("dataset_citation"):
            entries.append(meta["dataset_citation"])
        for number, entry in enumerate(entries, start=1):
            cursor.paragraph(f"[{number}] {entry}", size=REFERENCE_PT, space_after=1)


# ---------------------------------------------------------------- output

def find_soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    default = Path("C:/Program Files/LibreOffice/program/soffice.exe")
    return str(default) if default.is_file() else None


def convert_with_word(docx_path: Path) -> Path | None:
    """Export the PDF through Microsoft Word (Windows only); return the PDF path or None."""
    if sys.platform != "win32":
        return None
    pdf = docx_path.with_suffix(".pdf")
    script = ("$ErrorActionPreference = 'Stop'; $w = New-Object -ComObject Word.Application; $w.Visible = $false; "
              "try { $d = $w.Documents.Open($env:HAR_DOCX, $false, $true); "
              "$d.ExportAsFixedFormat($env:HAR_PDF, 17); $d.Close($false) } finally { $w.Quit() }")
    env = {**os.environ, "HAR_DOCX": str(docx_path), "HAR_PDF": str(pdf)}
    try:
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       check=True, capture_output=True, env=env, timeout=300)
    except (OSError, subprocess.SubprocessError):
        return None
    return pdf if pdf.is_file() else None


def convert_to_pdf(docx_path: Path) -> Path | None:
    """Convert with LibreOffice, or with Word on Windows; return the PDF path or None."""
    soffice = find_soffice()
    if soffice is None:
        return convert_with_word(docx_path)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True)
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.is_file() else None


def count_pages(pdf: Path) -> int:
    from pypdf import PdfReader

    return len(PdfReader(str(pdf)).pages)


def build(inputs: Inputs, template: Path, out_dir: Path, make_pdf: bool) -> Path:
    """Build the docx (and PDF if possible); drop the reference list if the PDF runs over 5 pages."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{OUTPUT_STEM}{'_DRAFT' if inputs.draft else ''}.docx"
    pdf = None
    for include_references in (True, False):
        inputs.problems.clear()
        ReportBuilder(inputs, template, include_references).build().save(str(docx_path))
        pdf = convert_to_pdf(docx_path) if make_pdf else None
        if pdf is None:
            break
        pages = count_pages(pdf) - 1  # excluding the cover page
        print(f"PDF: {pdf} ({pages} pages excluding the cover)")
        if pages <= 5 or not include_references:
            if not 3 <= pages <= 5:
                print(f"WARNING: Part B should be 3 to 5 pages excluding the cover; this is {pages}.")
            break
        print("Over 5 pages: rebuilding without the reference list (Part B does not require it).")
    if make_pdf and pdf is None:
        print("No PDF converter worked (LibreOffice or Word): open the docx in Word, check it and export the PDF.")
    return docx_path


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the ICT 4442 interim report (Part B) from repo files.")
    parser.add_argument("--root", type=Path, default=REPO_ROOT, help="tree to read inputs from (default: repo root)")
    parser.add_argument("--template", type=Path, default=TEMPLATE, help="official report template docx")
    parser.add_argument("--out-dir", type=Path, default=None, help="output folder (default: <root>/report/build)")
    parser.add_argument("--draft", action="store_true",
                        help="allow placeholders and missing inputs (writes ..._DRAFT.docx)")
    parser.add_argument("--no-pdf", action="store_true", help="skip PDF conversion")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> Path:
    """Build the report and return the docx path."""
    args = parse_args(argv)
    inputs = Inputs(root=args.root.resolve(), draft=args.draft)
    out_dir = args.out_dir or inputs.path("report", "build")
    docx_path = build(inputs, args.template, out_dir, make_pdf=not args.no_pdf)
    for message in inputs.problems:
        print(f"DRAFT WARNING: {message}")
    print(f"Report written to {docx_path}. Open it in Word and check every number against results/summary.csv.")
    return docx_path


if __name__ == "__main__":
    try:
        main()
    except BuildError as exc:
        sys.exit(f"build_interim_report: {exc}")
