"""Tests for report/build_interim_report.py, run against dummy inputs in a temporary folder.

The dummy files below exist only inside pytest's tmp_path; their numbers are placeholders,
never results.
"""

from __future__ import annotations

import importlib.util
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from types import ModuleType

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pytest  # noqa: E402
from docx import Document  # noqa: E402

from har.utils.paths import REPO_ROOT  # noqa: E402


def _load_builder() -> ModuleType:
    path = REPO_ROOT / "report" / "build_interim_report.py"
    spec = importlib.util.spec_from_file_location("build_interim_report", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


builder = _load_builder()

LIT_COUNTS = {"mayank": 4, "ishan": 5, "addyan": 4}


def _write(root: Path, rel: str, text: str) -> None:
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _png(root: Path, rel: str) -> None:
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(2, 2))
    ax.plot([0, 1], [0, 1])
    fig.savefig(path, dpi=50)
    plt.close(fig)


def _lit(member: str, n: int) -> str:
    rows = "\n".join(f"| Dummy{member}{i}, 2020 | method {i} | dataset {i} | result {i} | relevance {i} |"
                     for i in range(n))
    refs = "\n".join(f"- A. Dummy{member}{i}, \"Title {i},\" *Journal*, 2020." for i in range(n))
    return ("| Paper (Author, Year) | Method | Dataset | Key Result | Relevance to Project |\n"
            "|---|---|---|---|---|\n" f"{rows}\n\n## References\n{refs}\n")


def _set_meta(meta: str, key: str, value: str) -> str:
    """Replace the value of a top-level key in meta.yaml text."""
    return re.sub(rf"^{key}:.*$", f"{key}: {value}", meta, count=1, flags=re.M)


def _tasks(prefix: str, name: str) -> str:
    return (f"# Tasks: {name}\n\n| ID | Task | Status | Branch |\n|----|------|--------|--------|\n"
            f"| {prefix}.1 | First dummy task | Done | x/a |\n| {prefix}.2 | Second dummy task | Todo | x/b |\n")


@pytest.fixture()
def dummy_root(tmp_path: Path) -> Path:
    root = tmp_path / "dummy"
    meta = (REPO_ROOT / "report" / "interim" / "meta.yaml").read_text(encoding="utf-8")
    meta = _set_meta(meta, "team_no", '"99"')
    meta = _set_meta(meta, "repo_url", "https://github.com/example/dummy.git")
    _write(root, "report/interim/meta.yaml", meta)
    _write(root, "configs/team.yaml", (REPO_ROOT / "configs" / "team.yaml").read_text(encoding="utf-8"))
    _write(root, "report/interim/members.local.yaml", 'M1: "REG-1"\nM2: "REG-2"\nM3: "REG-3"\n')
    for member, n in LIT_COUNTS.items():
        _write(root, f"docs/lit/{member}.md", _lit(member, n))
    for prefix, member, name in (("M1", "mayank", "Mayank"), ("M2", "ishan", "Ishan"), ("M3", "addyan", "Addyan")):
        _write(root, f"docs/tasks/{member}.md", _tasks(prefix, name))
    sections = "report/interim/sections/"
    for name in ("lit_summary", "preprocessing_raw", "preprocessing_features", "evaluation_protocol", "observations"):
        _write(root, f"{sections}{name}.md", f"<!-- hidden comment -->\nDummy text for {name}.\n")
    for key in ("mlp", "cnn1d", "transformer"):
        _write(root, f"{sections}model_{key}.md",
               f'---\nmodel: {key}\nstatus: "Dummy status {key}"\nnotes: "Dummy notes {key}"\n---\nSummary of {key}.\n')
    _write(root, f"{sections}model_bilstm.md",
           '---\nmodel: bilstm\nstatus: "Dummy status bilstm"\nnotes: "Dummy notes bilstm"\n'
           'ablations:\n  gru: {status: "Dummy status gru", notes: "Dummy notes gru"}\n---\nSummary of bilstm and gru.\n')
    _write(root, f"{sections}risk_plan.md",
           "# Risk / Plan\n\n## Remaining work\n\n- Dummy bullet one\n- Dummy bullet two\n\n"
           "## Timeline to Final\n\n| Dates | Milestone | Owner |\n|---|---|---|\n| 1 Jan | Dummy milestone | All |\n")
    _write(root, "results/summary.csv",
           "model,run_by,val_acc,val_macro_f1,test_acc,test_macro_f1\n"
           "mlp,Dummy,0.7,0.6,0.5,0.5\ncnn1d,Dummy,0.8,0.7,0.6,0.6\nbilstm,Dummy,0.7,0.65,,\n")
    _png(root, "docs/figures/class_distribution.png")
    _png(root, "results/figures/confusion_test_cnn1d.png")
    _write(root, "docs/contribution_log.md",
           "# Contribution log\n\nHEAD: abcdef1234567\n\n| Member | Commits |\n|---|---|\n| Dummy | 1 |\n")
    return root


def _build(root: Path, *extra: str) -> Path:
    return builder.main(["--root", str(root), "--no-pdf", "--out-dir", str(root / "out"), *extra])


def _all_text(doc: Document) -> str:
    return "\n".join(t.text or "" for t in doc.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))


# ---------------------------------------------------------------- markdown helpers

def test_read_markdown_strips_front_matter_and_comments(tmp_path: Path) -> None:
    path = tmp_path / "x.md"
    path.write_text("---\nstatus: ok\n---\n<!-- note -->\nBody text.\n", encoding="utf-8")
    assert builder.read_markdown(path) == ({"status": "ok"}, "Body text.")


def test_markdown_blocks_and_tables() -> None:
    blocks = builder.markdown_blocks("## Head\n\nLine one\nline two.\n\n- a bullet\n  wrapped\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    assert [(b.kind, b.text) for b in blocks[:3]] == [("heading", "Head"), ("para", "Line one line two."),
                                                      ("bullet", "a bullet wrapped")]
    assert blocks[3].table == (["A", "B"], [["1", "2"]])


def test_preliminary_metric_format() -> None:
    row = {"test_acc": "0.9512", "test_macro_f1": "0.95", "val_macro_f1": "0.961"}
    assert builder.preliminary_metric(row) == "Test acc 95.1% / macro-F1 0.950 (val macro-F1 0.961)"
    assert builder.preliminary_metric(None) == "Not trained yet"
    assert builder.preliminary_metric({"val_macro_f1": "0.9", "test_acc": ""}).startswith("Test not evaluated")


def test_https_repo_url() -> None:
    assert builder.https_repo_url("git@github.com:o/r.git") == "https://github.com/o/r"
    assert builder.https_repo_url("https://user@github.com/o/r.git\n") == "https://github.com/o/r"


# ---------------------------------------------------------------- full build on dummy inputs

def test_build_fills_part_b_and_removes_parts_a_and_c(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    text = _all_text(doc)
    for gone in ("PART A", "PART C", "Synopsis/Interim/Final", "<NAME OF THE STUDENT", "MONTH AND YEAR",
                 "Interim report should be", "Minimum 8", "hidden comment"):
        assert gone not in text, gone
    for present in ("Interim report on", "INTERIM REPORT", "Mayank Kejariwal - REG-1", "Addyan Kumar - REG-3",
                    "October 2026", "Team 99:", "https://github.com/example/dummy", "Confirmed from Synopsis",
                    "Dummy text for preprocessing_raw.", "Dummy bullet two", "at commit abcdef1"):
        assert present in text, present
    # the page setup survives and the document does not end on an empty page-break paragraph
    assert doc.element.body[-1].tag.endswith("sectPr")
    assert doc.paragraphs[-1].text.strip()


def test_tables_are_filled_by_header(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    lit = builder.find_table(doc, builder.LIT_HEADERS)
    papers = [row.cells[0].text for row in lit.rows[1:]]
    assert len(papers) == sum(LIT_COUNTS.values()) == 13
    assert papers[0] == "Dummymayank0, 2020 [1]" and papers[-1] == "Dummyaddyan3, 2020 [13]"

    models = builder.find_table(doc, builder.MODEL_HEADERS)
    rows = {r.cells[0].text: [c.text for c in r.cells] for r in models.rows[1:]}
    assert len(rows) == 5
    assert rows["MLP (engineered features)"][1:] == [
        "Mayank Kejariwal", "Dummy status mlp", "Test acc 50.0% / macro-F1 0.500 (val macro-F1 0.600)", "Dummy notes mlp"]
    assert rows["BiGRU ablation (raw signal)"][2] == "Dummy status gru"
    assert rows["BiLSTM (raw signal)"][3].startswith("Test not evaluated")
    assert rows["Transformer encoder (raw signal)"][3] == "Not trained yet"

    contrib = builder.find_table(doc, builder.CONTRIB_HEADERS)
    assert [r.cells[2].text for r in contrib.rows[1:]] == [
        "M1.1 First dummy task", "M2.1 First dummy task", "M3.1 First dummy task"]
    assert [r.cells[3].text for r in contrib.rows[1:]] == ["", "", ""]  # signed by hand


def test_captions_figures_and_references(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    text = _all_text(doc)
    captions = re.findall(r"TABLE ([IVX]+)\.", text)
    assert captions == ["I", "II", "III", "IV", "V"]
    assert "TABLE V. TIMELINE TO FINAL" in text
    assert "Fig. 1. Windows per activity class" in text
    assert "Fig. 2. Test-set confusion matrix of the 1D-CNN (raw signal)" in text  # best val macro-F1
    template_images = len(Document(str(builder.TEMPLATE)).inline_shapes)  # cover logo
    assert len(doc.inline_shapes) == template_images + 2
    refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\] ", p.text)]
    assert len(refs) == 14 and refs[-1].startswith("[14] J. Reyes-Ortiz")


def test_final_build_refuses_placeholders_but_draft_builds(dummy_root: Path) -> None:
    meta = dummy_root / "report/interim/meta.yaml"
    meta.write_text(_set_meta(meta.read_text(encoding="utf-8"), "team_no", '"FILL_ME"'), encoding="utf-8")
    with pytest.raises(builder.BuildError, match="team_no"):
        _build(dummy_root)
    (dummy_root / "report/interim/members.local.yaml").unlink()
    (dummy_root / "report/interim/sections/observations.md").unlink()
    path = _build(dummy_root, "--draft")
    assert path.name.endswith("_DRAFT.docx")
    text = _all_text(Document(str(path)))
    assert "Team TBD:" in text and "Reg. No. TBD" in text


def test_final_build_refuses_missing_inputs(dummy_root: Path) -> None:
    (dummy_root / "report/interim/sections/observations.md").unlink()
    with pytest.raises(builder.BuildError, match="observations.md"):
        _build(dummy_root)


# ---------------------------------------------------------------- stale results table guard

def _run_folder(root: Path, rel_dir: str, override: bool = False) -> None:
    run_id = rel_dir.rsplit("/", 1)[-1]
    _write(root, f"results/{rel_dir}/metrics.json", json.dumps({"run_id": run_id, "protocol_override": override}))


def test_run_started_parses_the_run_id() -> None:
    assert builder.run_started(Path("20261004-194905_134ae8c")) == datetime(2026, 10, 4, 19, 49, 5)
    assert builder.run_started(Path("not-a-run")) is None


def test_final_build_refuses_a_model_with_runs_but_no_summary_row(dummy_root: Path) -> None:
    _run_folder(dummy_root, "gru/20000101-000000_abc1234")  # the dummy summary has no gru row
    with pytest.raises(builder.BuildError, match="no 'gru' row"):
        _build(dummy_root)


def test_final_build_refuses_a_summary_older_than_the_newest_run(dummy_root: Path) -> None:
    _run_folder(dummy_root, "mlp/20991231-235959_abc1234")  # started after summary.csv was written
    with pytest.raises(builder.BuildError, match="older than the newest run folder"):
        _build(dummy_root)


def test_current_summary_passes_and_smoke_or_override_runs_are_ignored(dummy_root: Path) -> None:
    _run_folder(dummy_root, "mlp/20000101-000000_abc1234")
    _run_folder(dummy_root, "_smoke/gru/20991231-235959_abc1234")
    _run_folder(dummy_root, "gru/20991231-235959_abc1234", override=True)
    assert _build(dummy_root).is_file()


def test_draft_build_only_warns_about_a_stale_summary(dummy_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _run_folder(dummy_root, "gru/20991231-235959_abc1234")
    _build(dummy_root, "--draft")
    out = capsys.readouterr().out
    assert "no 'gru' row" in out and "older than the newest run folder" in out


def test_mismatched_literature_references_are_rejected(dummy_root: Path) -> None:
    path = dummy_root / "docs/lit/ishan.md"
    path.write_text(path.read_text(encoding="utf-8").rsplit("\n- ", 1)[0] + "\n", encoding="utf-8")
    with pytest.raises(builder.BuildError, match="references"):
        _build(dummy_root)
