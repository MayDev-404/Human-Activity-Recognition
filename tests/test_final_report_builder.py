"""Tests for report/build_final_report.py (PHASE3_SPEC 6.6), run against dummy inputs in a temporary folder.

The dummy files below exist only inside pytest's tmp_path; their numbers are placeholders,
never results.
"""

from __future__ import annotations

import csv
import importlib.util
import json
import re
import sys
from pathlib import Path
from types import ModuleType

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pytest  # noqa: E402
import yaml  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from har.data.constants import ACTIVITY_NAMES  # noqa: E402
from har.eval.results import FINAL_CSV_COLUMNS  # noqa: E402
from har.utils.paths import REPO_ROOT  # noqa: E402


def _load_builder() -> ModuleType:
    path = REPO_ROOT / "report" / "build_final_report.py"
    spec = importlib.util.spec_from_file_location("build_final_report", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


builder = _load_builder()

MODELS = ["mlp", "cnn1d", "bilstm", "gru", "transformer"]
LABELS = {"mlp": "MLP (features)", "cnn1d": "1D-CNN", "bilstm": "BiLSTM", "gru": "BiGRU (ablation)",
          "transformer": "Transformer"}
DUMMY_F1 = {"mlp": 0.511, "cnn1d": 0.533, "bilstm": 0.522, "gru": 0.520, "transformer": 0.515}  # cnn1d best
INTRO_CITES = "[@anguita2013] and [@ucihar2013; @reyesortiz2016]"
LIT_KEYS = ["mahmud2020", "dirgova2022", "yang2015", "jiang2015", "ronao2016", "wang2019", "xia2020",
            "ordonez2016", "hammerla2016", "murad2017"]
SPECIAL = {  # extra text per section file, on top of filler words within the budget
    "1_introduction.md": f"Intro cites {INTRO_CITES}.",
    "2_literature.md": " ".join(f"[@{k}]" for k in LIT_KEYS) + " and again [@anguita2013].",
    "5_setup.md": "Validation subjects {{split:val_subjects}}; MLP config {{tuning:mlp:config_id}} in {{table:tuning}}.",
    "6a_results.md": ("CNN {{final:cnn1d:test_macro_f1}} (mean {{final:cnn1d:test_macro_f1.mean}}), "
                      "MLP accuracy {{final:mlp:test_acc}}, see {{table:final}} and {{fig:cost_tradeoff}}."),
    "6b_error_analysis.md": "Total errors {{analysis:errors.total}}, share {{analysis:errors.share}}.",
}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _write(root: Path, rel: str, text: str) -> None:
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _png(root: Path, rel: str) -> None:
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(2, 2))
    ax.plot([0, 1], [0, 1])
    fig.savefig(path, dpi=50)
    plt.close(fig)


def _filler(n: int) -> str:
    return " ".join(f"filler{i}" for i in range(n))


def _section_text(part) -> str:
    if part is builder.KEYWORDS:
        return "alpha, beta, gamma, delta, epsilon, zeta\n"
    special = SPECIAL.get(part.file, "")
    middle = sum(part.budget) // 2
    return f"# Heading dropped by the builder\n\n{special} {_filler(middle - builder.word_count(special))}\n"


def _all_parts():
    return [builder.ABSTRACT, builder.KEYWORDS, builder.ACKNOWLEDGEMENT] + [p for _, ps in builder.OUTLINE for p in ps]


def _csv(root: Path, rel: str, header: list[str], rows: list[dict]) -> None:
    path = root / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=header)
        writer.writeheader()
        writer.writerows(rows)


META = {
    "team_no": "99",
    "title": "Dummy Title of the Project",
    "month_year": "October 2026",
    "affiliation": "Dummy School, Dummy Institute",
    "models": [{"key": k, "label": LABELS[k], "owner": "M1"} for k in MODELS],
}


@pytest.fixture()
def dummy_root(tmp_path: Path) -> Path:
    root = tmp_path / "dummy"
    _write(root, "report/final/meta.yaml", yaml.safe_dump(META, sort_keys=False))
    for rel in ("configs/team.yaml", "configs/base.yaml", "configs/split.json", "report/final/references.yaml"):
        _write(root, rel, (REPO_ROOT / rel).read_text(encoding="utf-8"))
    _write(root, "report/interim/members.local.yaml",
           'M1: "REG-1"\nM2: "REG-2"\nM3: "REG-3"\n'
           'emails: {M1: "m1@example.com", M2: "m2@example.com", M3: "m3@example.com"}\n')
    for part in _all_parts():
        _write(root, f"report/final/sections/{part.file}", _section_text(part))
    for member in ("mayank", "ishan", "addyan"):
        _write(root, f"report/final/contributions/{member}.md",
               f'---\nmodel_task: "Dummy task of {member}"\n---\n<!-- note -->\nDummy contribution of {member}.\n')

    rows = []
    for i, model in enumerate(MODELS):
        row = {c: "" for c in FINAL_CSV_COLUMNS}
        row.update(model=model, run_by="Dummy", representation="raw", params=1000 + i, seeds="42 43 44",
                   val_acc_mean=0.55, val_acc_std=0.01, val_macro_f1_mean=0.544, val_macro_f1_std=0.003,
                   test_acc_mean=0.6123, test_acc_std=0.0034, test_macro_f1_mean=DUMMY_F1[model],
                   test_macro_f1_std=0.004, train_time_s_mean=12.34, train_time_s_std=1.2,
                   cpu_latency_bs1_ms_mean=1.5, cpu_latency_source="profile_dummy.csv (single process)")
        rows.append(row)
    _csv(root, "results/final_summary.csv", FINAL_CSV_COLUMNS, rows)
    header = ["config_id", "dropout", "width", "params", "val_acc", "val_macro_f1", "best_epoch", "run_id", "selected"]
    for model in MODELS:
        _csv(root, f"results/tuning/{model}/tuning_summary.csv", header, [
            {"config_id": f"c{i + 1}", "dropout": d, "width": w, "params": 100 * (i + 1), "val_acc": 0.5,
             "val_macro_f1": 0.5 + i / 100, "best_epoch": 3, "run_id": f"r{i}", "selected": str(i == 1).lower()}
            for i, (d, w) in enumerate([(0.1, 8), (0.1, 16), (0.2, 8), (0.2, 16)])])
    _csv(root, "results/analysis/per_class_f1.csv", ["model", "class", "f1_mean", "f1_std"],
         [{"model": m, "class": c, "f1_mean": 0.4, "f1_std": 0.01} for m in MODELS for c in ACTIVITY_NAMES])
    _write(root, "results/analysis/error_summary.json", json.dumps({"errors": {"total": 12, "share": 0.25}}))
    for rel in ("docs/figures/class_distribution.png", "results/figures/final_val_curves.png",
                "results/figures/cost_tradeoff.png", "results/analysis/per_class_f1.png",
                "results/analysis/error_overlap.png", "results/figures/confusion_test_cnn1d_final.png"):
        _png(root, rel)
    return root


def _build(root: Path, *extra: str) -> Path:
    return builder.main(["--root", str(root), "--no-pdf", "--out-dir", str(root / "out"), *extra])


def _all_text(doc: Document) -> str:
    return "\n".join(t.text or "" for t in doc.element.body.iter(W + "t"))


def _outline(doc: Document, level: int) -> list:
    found = []
    for paragraph in doc.paragraphs:
        lvl = paragraph._p.find(f"{W}pPr/{W}outlineLvl")
        if lvl is not None and lvl.get(qn("w:val")) == str(level):
            found.append(paragraph)
    return found


def _paragraph(doc: Document, prefix: str):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def _append(root: Path, file: str, text: str) -> None:
    path = root / "report/final/sections" / file
    path.write_text(path.read_text(encoding="utf-8") + text, encoding="utf-8")


# ---------------------------------------------------------------- helpers

def test_roman_and_float_numbers() -> None:
    assert [builder.roman(n) for n in (1, 4, 6, 8, 9, 14)] == ["I", "IV", "VI", "VIII", "IX", "XIV"]
    numbers = builder.float_numbers()
    assert [numbers[f"table:{k}"] for k in ("final", "tuning", "per_class")] == ["Table I", "Table II", "Table III"]
    assert numbers["fig:class_distribution"] == "Fig. 1" and numbers["fig:confusion_best"] == "Fig. 6"


def test_format_value() -> None:
    f1, acc = (".3f", 1, ""), (".1f", 100, "%")
    assert builder.format_value(0.9412, 0.0041, f1, None) == "0.941 ± 0.004"
    assert builder.format_value(0.9412, 0.0041, acc, None) == "94.1 ± 0.4%"
    assert builder.format_value(0.9412, 0.0041, f1, "mean") == "0.941"
    assert builder.format_value(0.9412, 0.0041, acc, "std") == "0.4%"
    assert builder.format_value(1.25, None, f1, None) == "1.250"   # single-process latency has no std
    assert builder.format_value(None, None, f1, None) is None


def test_word_count_ignores_citations_and_counts_placeholders_once() -> None:
    assert builder.word_count("Two **bold** words [@a; @b] and {{final:mlp:test_acc}}.") == 5


# ---------------------------------------------------------------- full build on dummy inputs

def test_build_fills_the_cover_and_removes_the_template_parts(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    text = _all_text(doc)
    for gone in ("PART A", "PART B", "PART C", "Synopsis/Interim/Final", "<NAME OF THE STUDENT", "MONTH AND YEAR",
                 "C.4 : Formatting Notes", "Fill in every field", "TEMPLATE", "hidden", "Heading dropped",
                 "<!--", "note -->"):
        assert gone not in text, gone
    for present in ("Final report on", "Mayank Kejariwal - REG-1", "Addyan Kumar - REG-3", "October 2026",
                    "Dummy Title of the Project"):
        assert present in text, present
    # the final sectPr is kept, set to A4, with no headers, footers or page numbers
    body = doc.element.body
    assert body[-1].tag == qn("w:sectPr")
    size = body[-1].find(qn("w:pgSz"))
    assert (size.get(qn("w:w")), size.get(qn("w:h"))) == ("11906", "16838")
    assert body[-1].find(qn("w:headerReference")) is None and body[-1].find(qn("w:footerReference")) is None


def test_title_block_fonts_and_body_text(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    title = [p for p in doc.paragraphs if p.text == "Dummy Title of the Project"][-1]  # the cover has one too
    assert title.runs[0].font.size.pt == 24
    authors = _paragraph(doc, "Mayank Kejariwal, Ishan Abhijit Saraf, Addyan Kumar")
    assert authors.runs[0].font.size.pt == 11
    affiliation = _paragraph(doc, "Dummy School")
    assert affiliation.runs[0].italic and affiliation.runs[0].font.size.pt == 10
    emails = _paragraph(doc, "m1@example.com")
    assert emails.text == "m1@example.com, m2@example.com, m3@example.com"
    assert emails.runs[0].font.size.pt == 9 and emails.runs[0].font.name == "Courier New"
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(f"{W}rPrDefault/{W}rPr/{W}rFonts")
    assert defaults.get(qn("w:ascii")) == "Times New Roman"
    abstract = _paragraph(doc, "Abstract: ")
    assert abstract.runs[0].bold and abstract.runs[0].italic
    assert builder.EM_DASH not in abstract.runs[0].text   # run-in label without an em dash
    keywords = _paragraph(doc, "Keywords: ")
    assert keywords.text == "Keywords: alpha, beta, gamma, delta, epsilon, zeta"
    body = _paragraph(doc, "Intro cites")
    assert body.alignment == 3 and all(r.font.size.pt == 10 for r in body.runs)  # justified, 10 pt


def test_headings_are_numbered_and_formatted(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    level1 = _outline(doc, 0)
    assert [p.text for p in level1] == [
        "I. Introduction", "II. Literature Review", "III. Objectives", "IV. Methodology", "V. Experimental Setup",
        "VI. Results and Discussion", "VII. Conclusion", "VIII. Future Work", "Acknowledgement", "References",
        "Individual Contribution Statement"]
    assert all(p.alignment == 1 and p.runs[0].font.small_caps for p in level1)  # centred small caps
    level2 = _outline(doc, 1)
    assert [p.text for p in level2] == [
        "A. Model 1: MLP", "B. Model 2: 1D-CNN", "C. Model 3: BiLSTM with GRU Ablation",
        "D. Model 4: Transformer Encoder", "A. Comparison of the Models", "B. Error Analysis"]
    assert all(p.alignment == 0 and p.runs[0].italic for p in level2)  # italic, left


def test_captions_tables_above_figures_below(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    body = list(doc.element.body.iterchildren())
    tables = [i for i, el in enumerate(body) if el.tag == qn("w:tbl")]
    captions = [builder.interim.element_text(body[i - 1]) for i in tables[:-1]]  # the last one is C.3
    assert [c.split(".")[0] for c in captions] == ["TABLE I", "TABLE II", "TABLE III"]
    caption = _paragraph(doc, "TABLE I. ")
    assert caption.runs[0].font.small_caps and caption.runs[0].font.size.pt == 8 and caption.alignment == 1
    figures = [p for p in doc.paragraphs if re.match(r"^Fig\. \d+\. ", p.text)]
    assert [p.text.split(".")[1].strip() for p in figures] == ["1", "2", "3", "4", "5", "6"]
    for paragraph in figures:
        assert paragraph.runs[0].font.size.pt == 8
        previous = paragraph._p.getprevious()
        assert previous.findall(".//" + qn("w:drawing")), paragraph.text   # caption sits below its figure
    assert "Fig. 6. Test confusion matrix of the 1D-CNN, the model with the best mean test macro-F1" in _all_text(doc)


def test_citations_are_numbered_by_first_citation(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    assert _paragraph(doc, "Intro cites").text.startswith("Intro cites [1] and [2], [3].")
    literature = _paragraph(doc, "[4] [5]").text
    assert literature.startswith("[4] [5] [6] [7] [8] [9] [10] [11] [12] [13] and again [1].")
    references = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\] ", p.text) and p.runs[0].font.size.pt == 8]
    assert len(references) == 13
    assert references[0].startswith("[1] D. Anguita") and references[1].startswith("[2] J. Reyes-Ortiz")
    assert all(p.text != "[14]" for p in doc.paragraphs)


def test_placeholders_are_resolved(dummy_root: Path) -> None:
    text = _all_text(Document(str(_build(dummy_root))))
    assert "CNN 0.533 ± 0.004 (mean 0.533), MLP accuracy 61.2 ± 0.3%, see Table I and Fig. 3." in text
    assert "Validation subjects 3, 15, 19, 22, 28, 29; MLP config c2 in Table II." in text
    assert "Total errors 12, share 0.250." in text
    assert "{{" not in text and "[@" not in text and "MISSING" not in text


def test_generated_tables(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    final, tuning, per_class = ([[c.text for c in r.cells] for r in t.rows] for t in doc.tables[:3])
    assert final[0] == ["Model", "Params", "Val Macro-F1", "Test Acc", "Test Macro-F1", "Train time (s)",
                        "CPU latency (ms)"]
    assert final[2] == ["1D-CNN", "1,001", "0.544 ± 0.003", "61.2 ± 0.3%", "0.533 ± 0.004", "12.3 ± 1.2", "1.500"]
    assert [r[0] for r in final[1:]] == list(LABELS.values())
    assert tuning[1] == ["MLP (features)", "dropout {0.1, 0.2} × width {8, 16}", "c2: dropout 0.1, width 16",
                         "200", "0.510"]
    assert per_class[0] == ["Class"] + list(LABELS.values())
    assert [r[0] for r in per_class[1:]] == ["Walking", "Walking upstairs", "Walking downstairs", "Sitting",
                                             "Standing", "Laying"]
    assert per_class[1][1] == "0.400 ± 0.010"
    assert "CPU latency at batch size 1 on one thread: profile_dummy.csv (single process)." in _all_text(doc)


def test_contribution_page_is_last_with_the_c3_table(dummy_root: Path) -> None:
    doc = Document(str(_build(dummy_root)))
    c3 = doc.tables[-1]
    assert [c.text.strip() for c in c3.rows[0].cells][:2] == ["Member Name", "Reg. No."]
    assert [[c.text for c in r.cells] for r in c3.rows[1:]] == [
        ["Mayank Kejariwal", "REG-1", "Dummy task of mayank", "Dummy contribution of mayank.", ""],
        ["Ishan Abhijit Saraf", "REG-2", "Dummy task of ishan", "Dummy contribution of ishan.", ""],
        ["Addyan Kumar", "REG-3", "Dummy task of addyan", "Dummy contribution of addyan.", ""]]  # 4th row deleted
    heading = _paragraph(doc, "Individual Contribution Statement")._p
    breaks = heading.getprevious().findall(".//" + qn("w:br"))
    assert breaks and breaks[0].get(qn("w:type")) == "page"   # C.3 starts a new page
    assert "Team 99: Dummy Title of the Project" in _all_text(doc)
    assert doc.element.body[-2] is c3._tbl                      # nothing after it but the sectPr


def test_core_properties_name_the_members(dummy_root: Path) -> None:
    props = Document(str(_build(dummy_root))).core_properties
    assert props.title == "Dummy Title of the Project"
    assert props.author == "Mayank Kejariwal; Ishan Abhijit Saraf; Addyan Kumar"


# ---------------------------------------------------------------- guards

def test_unresolved_placeholder_fails_and_draft_marks_it(dummy_root: Path) -> None:
    _append(dummy_root, "6a_results.md", "\nBad {{final:cnn1d:no_such_field}} and {{tuning:nope:x}}.\n")
    with pytest.raises(builder.BuildError, match="cannot resolve"):
        _build(dummy_root)
    text = _all_text(Document(str(_build(dummy_root, "--draft"))))
    assert "Bad [MISSING: final:cnn1d:no_such_field] and [MISSING: tuning:nope:x]." in text


def test_unknown_citation_key_fails_and_draft_marks_it(dummy_root: Path) -> None:
    _append(dummy_root, "7_conclusion.md", "\nSee [@nobody2099].\n")
    with pytest.raises(builder.BuildError, match="unknown citation key @nobody2099"):
        _build(dummy_root)
    assert "See [MISSING: @nobody2099]." in _all_text(Document(str(_build(dummy_root, "--draft"))))


def test_too_few_papers_in_section_ii_fail(dummy_root: Path) -> None:
    (dummy_root / "report/final/sections/2_literature.md").write_text(
        " ".join(f"[@{k}]" for k in LIT_KEYS[:9]) + " " + _filler(480), encoding="utf-8")
    with pytest.raises(builder.BuildError, match="Section II .* cites 9 papers"):
        _build(dummy_root)


def test_too_few_references_fail(dummy_root: Path) -> None:
    (dummy_root / "report/final/sections/1_introduction.md").write_text(_filler(350), encoding="utf-8")
    with pytest.raises(builder.BuildError, match="only 11 references are cited"):
        _build(dummy_root)


def test_em_dash_fails(dummy_root: Path) -> None:
    _append(dummy_root, "3_objectives.md", f"\nOne {builder.EM_DASH} two.\n")
    with pytest.raises(builder.BuildError, match="em dash"):
        _build(dummy_root)


@pytest.mark.parametrize("rel", ["report/final/sections/5_setup.md", "report/final/contributions/ishan.md",
                                 "results/final_summary.csv", "results/figures/cost_tradeoff.png"])
def test_missing_inputs_fail_and_draft_marks_them(dummy_root: Path, rel: str) -> None:
    (dummy_root / rel).unlink()
    with pytest.raises(builder.BuildError, match=re.escape(Path(rel).name)):
        _build(dummy_root)
    text = _all_text(Document(str(_build(dummy_root, "--draft"))))
    assert "[MISSING: " in text


@pytest.mark.parametrize("key, value", [("team_no", "TBD"), ("month_year", "<MONTH AND YEAR>"), ("affiliation", "")])
def test_meta_placeholders_fail(dummy_root: Path, key: str, value: str) -> None:
    meta = {**META, key: value}
    _write(dummy_root, "report/final/meta.yaml", yaml.safe_dump(meta, sort_keys=False))
    with pytest.raises(builder.BuildError, match=key):
        _build(dummy_root)


def test_missing_email_fails(dummy_root: Path) -> None:
    _write(dummy_root, "report/interim/members.local.yaml", 'M1: "REG-1"\nM2: "REG-2"\nM3: "REG-3"\n')
    with pytest.raises(builder.BuildError, match="no email for M1"):
        _build(dummy_root)
    assert "[MISSING: email of M1]" in _all_text(Document(str(_build(dummy_root, "--draft"))))


def test_budget_warnings_do_not_fail(dummy_root: Path, capsys: pytest.CaptureFixture[str]) -> None:
    (dummy_root / "report/final/sections/3_objectives.md").write_text("Too short.", encoding="utf-8")
    _build(dummy_root)
    out = capsys.readouterr().out
    assert "WARNING: report" in out and "3_objectives.md: 2 words, budget 80 to 140" in out
    assert "1_introduction.md" not in out   # within budget


def test_section_tables_are_rejected(dummy_root: Path) -> None:
    _append(dummy_root, "6a_results.md", "\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    with pytest.raises(builder.BuildError, match="markdown table"):
        _build(dummy_root)


# ---------------------------------------------------------------- the repository itself

def test_committed_meta_has_no_placeholders() -> None:
    meta = yaml.safe_load((REPO_ROOT / "report/final/meta.yaml").read_text(encoding="utf-8"))
    assert builder.interim.placeholder_fields(meta) == []
    assert [m["key"] for m in meta["models"]] == MODELS


def test_draft_build_of_the_repository_runs(tmp_path: Path) -> None:
    """The repository builds as a draft at any stage; missing inputs become [MISSING: ...] markers."""
    path = builder.main(["--draft", "--no-pdf", "--out-dir", str(tmp_path)])
    assert path.is_file() and path.name == "ICT4442_Final_Report_HAR_DRAFT.docx"
